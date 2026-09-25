"""Refinance & Capital Events V1 Stage 3 -- the Refinance & Capital Structure
Audit workbook (``docs/architecture/REFINANCE_CAPITAL_EVENTS_V1.md`` 25.1).

1. **Eligibility.** Produced only for the saved analysis the analyst ran, and
   only when that analysis carries an executed refinance. Every other state is
   a typed refusal; nothing is written.
2. **Conventions.** The accepted export conventions: protected sheets,
   pessimistic formula caches, no macros, no volatile functions, no external
   links, a frozen Anchor Results sheet and a Checks sheet.
3. **Reconciliation (opt-in native Excel).** Recalculated by Excel, every
   check passes for the F5 distribution, the F6 contribution, a positive-rate
   amortizing loan and the full-member two-loan refinance -- and a tampered
   Anchor figure fails its check, so the Checks sheet cannot pass vacuously.
"""

from __future__ import annotations

import io
import re
import sqlite3
import zipfile
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import _refinance_v1_fixtures as rf  # type: ignore[import-not-found]
import _refinance_v1_stage_2_fixtures as fx  # type: ignore[import-not-found]
import _refinance_v1_stage_3_fixtures as s3  # type: ignore[import-not-found]
from _p7_6_fixtures import detailed_deal, lease_level_deal  # type: ignore[import-not-found]
from anchor.deals import store
from anchor.exports.refinance import refinance_audit_filename
from anchor.exports.refinance.audit import CONTRACT_VERSION
from excel_native_recalc import native_recalc_enabled, recalculate_with_excel  # type: ignore[import-not-found]

SHEETS = [
    "Summary",
    "Inputs",
    "Retiring Debt",
    "Sizing",
    "Bridge",
    "Replacement Loan",
    "Common Equity",
    "Anchor Results",
    "Checks",
    "Audit Metadata",
]
VOLATILE = ("NOW(", "TODAY(", "RAND(", "RANDBETWEEN(", "OFFSET(", "INDIRECT(", "INFO(", "CELL(")
PASSING = ("Pass", "Pass (both unavailable)")
STATUS_COLUMN = 5


@pytest.fixture
def db(tmp_path: Path) -> Path:
    return tmp_path / "anchor.db"


@pytest.fixture
def client(db: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    return s3.client(db, monkeypatch)


def _export(client: TestClient, investment_id: str, strategy_id: str = "base") -> Any:
    fingerprint = s3.analysis(client, investment_id, strategy_id)["structured_source_fingerprint"]
    return s3.audit(client, investment_id, fingerprint, strategy_id=strategy_id)


def _code(response: Any) -> str:
    """The refusal's code; an export that was produced instead fails here."""

    assert response.status_code != 200, "the workbook was exported"
    return response.json()["detail"]["code"]


def _dump(db: Path) -> list[str]:
    connection = sqlite3.connect(db)
    try:
        return list(connection.iterdump())
    finally:
        connection.close()


# =============================================================================
# 1. Eligibility
# =============================================================================


def test_f5_downloads_the_audit_with_every_sheet(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db, name="Harbor / Point")
    response = _export(client, investment_id)

    assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "no-store"
    assert 'filename="Harbor - Point - Refinance & Capital Structure Audit.xlsx"' in response.headers[
        "content-disposition"
    ]
    assert refinance_audit_filename("Harbor") == "Harbor - Refinance & Capital Structure Audit.xlsx"
    assert load_workbook(io.BytesIO(response.content)).sheetnames == SHEETS


def test_the_audit_requires_the_analysis_on_screen(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db)
    s3.analysis(client, investment_id)

    missing = s3.audit(client, investment_id, None)
    stale = s3.audit(client, investment_id, "0" * 64)
    assert (missing.status_code, _code(missing)) == (409, "analysis_missing")
    assert (stale.status_code, _code(stale)) == (409, "analysis_stale")


def test_a_structure_change_after_the_analysis_stales_the_audit(db: Path, client: TestClient) -> None:
    deal, investment_id = s3.refinance_deal(db)
    fingerprint = s3.analysis(client, investment_id)["structured_source_fingerprint"]
    store.set_deal_capital_structure(deal.id, s3.f6_structure(deal.id), db_path=db)

    assert _code(s3.audit(client, investment_id, fingerprint)) == "analysis_stale"


def test_a_structure_without_a_refinance_is_refused(db: Path, client: TestClient) -> None:
    deal = fx.base_deal(db, name="Plain")
    investment_id, _ = store.set_deal_capital_structure(deal.id, fx.closing_mezz_only(deal.id), db_path=db)
    assert investment_id is not None

    assert _code(_export(client, investment_id)) == "no_refinance"


def test_an_unexecuted_refinance_is_refused_never_exported_partially(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db, s3.f4_structure)
    response = _export(client, investment_id)

    assert _code(response) == "refinance_unavailable"
    assert fx.EVENT_ID not in response.json()["detail"]["message"]


def test_an_unknown_investment_or_strategy_is_not_found(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db)

    assert _code(s3.audit(client, "nope", "0" * 64)) == "investment_not_found"
    assert _code(s3.audit(client, investment_id, "0" * 64, strategy_id="nope")) == "investment_not_found"


def test_exporting_writes_nothing(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db)
    fingerprint = s3.analysis(client, investment_id)["structured_source_fingerprint"]
    before = _dump(db)

    for _ in range(2):
        assert s3.audit(client, investment_id, fingerprint).status_code == 200
    assert _dump(db) == before


# =============================================================================
# 2. Conventions
# =============================================================================


def _formulas(content: bytes) -> list[str]:
    book = load_workbook(io.BytesIO(content))
    return [
        cell.value
        for sheet in book.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]


def test_every_sheet_is_protected_and_the_package_has_no_macro_or_link(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db)
    content = _export(client, investment_id).content
    book = load_workbook(io.BytesIO(content))

    assert all(sheet.protection.sheet for sheet in book.worksheets)
    parts = s3.workbook_parts(content)
    assert not any("vbaProject" in part or "externalLink" in part for part in parts)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        content_types = archive.read("[Content_Types].xml").decode()
    assert "macroEnabled" not in content_types


def test_no_formula_is_volatile_or_reaches_outside_the_workbook(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db)
    formulas = _formulas(_export(client, investment_id).content)

    assert formulas
    for formula in formulas:
        upper = formula.upper()
        assert not any(name in upper for name in VOLATILE), formula
        assert not re.search(r"\[[^\]]*\]", formula), formula


def test_formula_caches_are_pessimistic(db: Path, client: TestClient) -> None:
    """Opened without recalculation, the workbook never claims a pass."""

    _, investment_id = s3.refinance_deal(db)
    content = _export(client, investment_id).content
    cached = load_workbook(io.BytesIO(content), data_only=True)

    statuses = [
        row[STATUS_COLUMN]
        for row in cached["Checks"].iter_rows(values_only=True)
        if len(row) > STATUS_COLUMN and isinstance(row[STATUS_COLUMN], str)
    ]
    assert not any(status in PASSING for status in statuses)
    assert "Not recalculated" in [row[1] for row in cached["Checks"].iter_rows(values_only=True) if row[0] == "Excel formulas recalculated"]


def test_anchor_results_freezes_the_engine_figures(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db)
    book = load_workbook(io.BytesIO(_export(client, investment_id).content))
    values = [cell.value for row in book["Anchor Results"].iter_rows() for cell in row]

    # F5: N = 8,000,000 - 5,520,000 - 80,000 - 40,000.
    for figure in (8_000_000, 5_520_000, 2_360_000):
        assert any(isinstance(value, (int, float)) and value == pytest.approx(figure) for value in values), figure
    metadata = {row[0]: row[1] for row in book["Audit Metadata"].iter_rows(values_only=True) if row[0]}
    assert CONTRACT_VERSION in metadata.values()


def test_no_internal_identity_is_printed(db: Path, client: TestClient) -> None:
    deal, investment_id = s3.refinance_deal(db)
    book = load_workbook(io.BytesIO(_export(client, investment_id).content))
    printed = [
        cell.value
        for sheet in book.worksheets
        if sheet.title != "Audit Metadata"
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and not cell.value.startswith("=")
    ]

    for identity in (fx.EVENT_ID, fx.REPLACEMENT_ID, deal.id, investment_id, "legacy_acquisition_loan"):
        assert not any(identity in value for value in printed), identity


@pytest.mark.parametrize("mode", ["detailed", "lease_level"])
def test_detailed_and_lease_level_deals_export_the_audit(mode: str, db: Path, client: TestClient) -> None:
    _, investment_id = CASES[mode](db)
    response = _export(client, investment_id)

    assert response.status_code == 200, response.text
    assert load_workbook(io.BytesIO(response.content)).sheetnames == SHEETS


def test_a_partnership_adds_the_partners_sheet(db: Path, client: TestClient) -> None:
    from _p7_9_fixtures import f1_terms  # type: ignore[import-not-found]

    deal, investment_id = s3.refinance_deal(db)
    assert client.put(f"/deals/{deal.id}/partnership", json={"partnership": s3.api_module._wire(f1_terms())}).status_code == 200
    names = load_workbook(io.BytesIO(_export(client, investment_id).content)).sheetnames

    assert names == [*SHEETS[:7], "Partners", *SHEETS[7:]]


# =============================================================================
# 3. Native Excel reconciliation (opt-in)
# =============================================================================


def _positive_rate(db: Path) -> tuple[Any, str]:
    deal = fx.base_deal(db, name="Positive rate", interest_rate=0.06)
    structure = fx.evented(
        deal.id,
        replacement_fees=(rf.lender_fee(80_000.0),),
        dscr=1.25,
        fixed=9_000_000.0,
        costs=(rf.third_party(40_000.0), rf.exit_fee(25_000.0, fx.legacy(deal.id))),
    )
    investment_id, _ = store.set_deal_capital_structure(deal.id, structure, db_path=db)
    assert investment_id is not None
    return deal, investment_id


def _full_member(db: Path) -> tuple[Any, str]:
    deal = fx.base_deal(db, name="Full member", interest_rate=0.055)
    investment_id, _ = store.set_deal_capital_structure(deal.id, fx.full_member_structure(deal.id), db_path=db)
    assert investment_id is not None
    s3.with_value_timepoint(db, investment_id, deal.id)
    return deal, investment_id


def _mode(maker: Any) -> Any:
    def build(db: Path) -> tuple[Any, str]:
        deal = maker(db, name=maker.__name__)
        investment_id, _ = store.set_deal_capital_structure(deal.id, s3.f5_structure(deal.id), db_path=db)
        assert investment_id is not None
        return deal, investment_id

    return build


CASES = {
    "detailed": _mode(detailed_deal),
    "lease_level": _mode(lease_level_deal),
    "f5": lambda db: s3.refinance_deal(db, s3.f5_structure),
    "f6": lambda db: s3.refinance_deal(db, s3.f6_structure),
    "positive_rate": _positive_rate,
    "full_member": _full_member,
}


def _statuses(path: Path) -> list[Any]:
    sheet = load_workbook(path, data_only=True)["Checks"]
    return [
        row[:STATUS_COLUMN + 1]
        for row in sheet.iter_rows(values_only=True)
        if len(row) > STATUS_COLUMN
        and isinstance(row[STATUS_COLUMN], str)
        and (row[STATUS_COLUMN].startswith(("Pass", "FAIL", "Missing", "Excel error", "Not like-for-like")))
    ]


def _recalculated(tmp_path: Path, name: str, content: bytes) -> Path:
    source = tmp_path / f"{name}.xlsx"
    target = tmp_path / f"{name}-recalculated.xlsx"
    source.write_bytes(content)
    recalculate_with_excel([(source, target)], tmp_path)
    return target


native = pytest.mark.skipif(not native_recalc_enabled(), reason="set ANCHOR_EXCEL_NATIVE_RECALC=1 to run Excel")


@native
@pytest.mark.parametrize("case", sorted(CASES))
def test_native_excel_reconciles_every_check(case: str, db: Path, client: TestClient, tmp_path: Path) -> None:
    _, investment_id = CASES[case](db)
    response = _export(client, investment_id)
    assert response.status_code == 200, response.text
    statuses = _statuses(_recalculated(tmp_path, case, response.content))

    assert len(statuses) > 40
    assert [row for row in statuses if row[STATUS_COLUMN] not in PASSING] == []


@native
def test_native_excel_fails_a_tampered_anchor_figure(db: Path, client: TestClient, tmp_path: Path) -> None:
    """Mutation proof: move one frozen Anchor figure by $1 and its check must
    fail after Excel recalculates."""

    _, investment_id = s3.refinance_deal(db)
    content = _export(client, investment_id).content
    book = load_workbook(io.BytesIO(content))
    sheet = book["Anchor Results"]
    sheet.protection.sheet = False
    (target,) = [
        cell
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, (int, float)) and cell.value == pytest.approx(2_360_000)
    ][:1]
    target.value = target.value + 1
    tampered = io.BytesIO()
    book.save(tampered)

    statuses = _statuses(_recalculated(tmp_path, "tampered", tampered.getvalue()))
    assert any(row[STATUS_COLUMN] == "FAIL" for row in statuses)


def test_every_refusal_code_has_its_own_status() -> None:
    """A new typed refusal cannot reach the route without an HTTP status."""

    from anchor.api import _REFINANCE_AUDIT_STATUS
    from anchor.exports.refinance.source import RefinanceAuditRefusalCode

    assert set(_REFINANCE_AUDIT_STATUS) == {code.value for code in RefinanceAuditRefusalCode}
