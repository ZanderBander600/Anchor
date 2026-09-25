"""Refinance & Capital Events V1 Stage 3 -- F23, primary-view semantics.

``docs/architecture/REFINANCE_CAPITAL_EVENTS_V1.md`` Section 18.2 (F23), proving
Section 12.5 end to end on one saved F5 Investment: the workspace's analysis,
the Decision Matrix, a newly published memo and report, and both exports.
Common Equity (or Partner) returns are the headline everywhere; the
acquisition-loan levered figures appear only under the acquisition-financing
reference label; an unavailable refinance (F11 shape) shows its state as
primary, never the acquisition figures. The frontend half -- both namespaces
named in a mixed matrix, no frontend arithmetic -- is proved by the Stage 3
vitest suite.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import _refinance_v1_stage_3_fixtures as s3  # type: ignore[import-not-found]
from _p7_9_fixtures import f1_terms  # type: ignore[import-not-found]
from anchor.deals import store
from anchor.deals.structured_variants import IMPLICIT_COMMON_EQUITY_ID

PRIMARY = "common_equity_after_capital_structure"
REFERENCE_NAMESPACE = "acquisition_financing_reference"


@pytest.fixture
def db(tmp_path: Path) -> Path:
    return tmp_path / "anchor.db"


@pytest.fixture
def client(db: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    return s3.client(db, monkeypatch)


def test_f23_the_workspace_analysis_states_the_primary_view(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db)
    primary = s3.analysis(client, investment_id)["primary_return"]

    assert primary["primary_equity_namespace"] == PRIMARY
    assert primary["reference_namespace"] == REFERENCE_NAMESPACE
    assert primary["reference_excludes_capital_events"] is True
    assert primary["status"] == "available"


def test_f23_the_matrix_names_the_reference_and_offers_common_equity(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db)
    presence = client.get(f"/investments/{investment_id}/capital-event-presence").json()
    common = client.post(f"/investments/{investment_id}/position-decision-matrix/{IMPLICIT_COMMON_EQUITY_ID}")

    assert "levered_irr" in presence["acquisition_financing_metrics"]
    assert "equity_multiple" in presence["acquisition_financing_metrics"]
    assert common.status_code == 200
    (cell,) = common.json()["matrix"]["cells"]
    assert {metric["metric"]: metric["value"] for metric in cell["metrics"]}["equity_multiple"] == pytest.approx(2.295)


def test_f23_a_published_report_leads_with_common_equity(db: Path) -> None:
    _, investment_id = s3.refinance_deal(db)
    version = s3.publish(db, investment_id)
    artifact = store.get_memo_version_artifact(investment_id, version.version_id, db_path=db)
    assert artifact is not None
    package = artifact.package()
    labels = [metric.label for metric in package.key_metrics]

    assert labels.index("Common Equity IRR") < labels.index("Cash Returned to Common Equity")
    assert "Levered IRR" not in labels
    returns = {metric.label: metric for metric in s3.section(package, "Returns").metrics}
    assert returns["Levered IRR"].note == s3.REFERENCE


def test_f23_a_partner_memo_leads_with_the_partner_return(db: Path, client: TestClient) -> None:
    deal, investment_id = s3.refinance_deal(db)
    assert client.put(
        f"/deals/{deal.id}/partnership", json={"partnership": s3.api_module._wire(f1_terms())}
    ).status_code == 200
    package = s3.preview(db, investment_id, s3.partner_cell("lp"))
    labels = [metric.label for metric in package.key_metrics]

    assert "Partner IRR – LP" in labels and "Common Equity IRR" not in labels
    assert s3.metric(package, "Partner IRR – LP").note == "Refinance-adjusted"
    assert s3.refusal_codes(db, investment_id) == []


def test_f23_both_exports_carry_the_primary_view(db: Path, client: TestClient) -> None:
    deal, investment_id = s3.refinance_deal(db)
    s3.analysed(db, deal)
    quick = load_workbook(io.BytesIO(client.get(f"/deals/{deal.id}/exports/quick-underwrite.xlsx").content))
    labels = [
        cell.value for row in quick["Summary"].iter_rows() for cell in row if isinstance(cell.value, str)
    ]
    fingerprint = s3.analysis(client, investment_id)["structured_source_fingerprint"]
    audit = s3.audit(client, investment_id, fingerprint)

    assert any(s3.REFERENCE in label for label in labels)
    assert audit.status_code == 200
    summary = [
        cell.value
        for row in load_workbook(io.BytesIO(audit.content))["Summary"].iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert any("Common Equity" in value for value in summary)


def test_f23_an_unavailable_refinance_is_primary_never_the_acquisition_figures(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db, s3.f4_structure)
    primary = s3.analysis(client, investment_id)["primary_return"]
    package = s3.preview(db, investment_id)

    assert primary["status"] != "available"
    assert primary["unavailable_message"]
    headline = s3.metric(package, "Common Equity IRR")
    assert headline.value is None and headline.unavailable is not None
    assert "Levered IRR" not in [metric.label for metric in package.key_metrics]
