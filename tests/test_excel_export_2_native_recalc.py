"""Excel Export 2 -- what a real spreadsheet engine calculates.

Opt-in (``ANCHOR_EXCEL_NATIVE_RECALC=1``), because it drives desktop Excel.
Everything else about the workbook is proved without one in
``test_excel_export_2_detailed_audit.py``; what needs Excel is the only thing
that cannot be faked: the *values* the formulas produce.

Three kinds of evidence here:

1. **Reconciliation.** Every golden case recalculates and every check passes,
   with no Excel error anywhere and every formula still a formula afterwards.
2. **Independent read-back.** Excel's own operating cells are compared with
   Anchor's engine in Python, without going through the Checks sheet -- so a
   Checks row that silently compared a cell with itself could not hide.
3. **Mutation.** Each invariant is broken in turn and the row that guards it
   must fail while an unrelated row still passes. A check that cannot fail is
   not a check.
"""

from __future__ import annotations

import dataclasses
from pathlib import Path

import pytest

from excel_detailed_export_golden_cases import (
    CASES_BY_NAME,
    DETAILED_GOLDEN_CASES,
    analyze,
    build,
    check_rows,
    load,
    number_at,
    operating,
    row_of,
    source_for,
    status_block,
    terms,
)
from excel_export_golden_cases import blank_cell, read_formula, replace_formula, set_number
from excel_native_recalc import native_recalc_enabled, recalculate_with_excel

pytestmark = pytest.mark.skipif(
    not native_recalc_enabled(),
    reason="set ANCHOR_EXCEL_NATIVE_RECALC=1 to recalculate in desktop Excel",
)

OPERATING = "Operating Projection"
CHECKS = "Checks"


def _recalculate(tmp_path: Path, workbooks: dict[str, bytes]) -> dict[str, bytes]:
    """Recalculate each named workbook in one Excel session."""

    pairs = []
    for name, data in workbooks.items():
        source = tmp_path / f"{name}.xlsx"
        source.write_bytes(data)
        pairs.append((source, tmp_path / f"{name}.recalc.xlsx"))
    recalculate_with_excel(pairs, tmp_path)
    return {
        name: destination.read_bytes()
        for name, (_, destination) in zip(workbooks, pairs)
    }


def _statuses(data: bytes) -> dict[str, str]:
    return {metric: row[3] for metric, row in check_rows(load(data, values=True)).items()}


# =============================================================================
# 1. Every golden case reconciles in Excel
# =============================================================================


@pytest.fixture(scope="module")
def recalculated(tmp_path_factory: pytest.TempPathFactory) -> dict[str, bytes]:
    """Every golden case, recalculated once, in a single Excel session."""

    work = tmp_path_factory.mktemp("detailed-native")
    return _recalculate(work, {case.name: build(case) for case in DETAILED_GOLDEN_CASES})


@pytest.mark.parametrize("case", DETAILED_GOLDEN_CASES, ids=lambda case: case.name)
def test_every_check_passes_after_a_real_recalculation(case, recalculated) -> None:  # noqa: ANN001
    data = recalculated[case.name]
    block = status_block(load(data, values=True))
    assert block["Excel formulas recalculated"] == "Yes"
    assert block["Anchor reconciliation available"] == "Yes"
    assert block["Workbook modified since export"] == "No"
    assert block["Original inputs unchanged"] == "Yes"

    statuses = _statuses(data)
    failures = {
        metric: status
        for metric, status in statuses.items()
        if not status.startswith("Pass") and metric not in case.expected_failures
    }
    assert failures == {}, failures
    assert block["Not passed"] == len(case.expected_failures)
    assert block["Passed"] == len(statuses) - len(case.expected_failures)


@pytest.mark.parametrize("case", DETAILED_GOLDEN_CASES, ids=lambda case: case.name)
def test_no_cell_holds_an_excel_error(case, recalculated) -> None:  # noqa: ANN001
    wb = load(recalculated[case.name], values=True)
    errors = []
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    errors.append((name, cell.coordinate, cell.value))
    assert errors == []


@pytest.mark.parametrize("case", DETAILED_GOLDEN_CASES, ids=lambda case: case.name)
def test_formulas_are_still_formulas_after_excel_saves(case, recalculated) -> None:  # noqa: ANN001
    original = load(build(case))
    saved = load(recalculated[case.name])
    for name in original.sheetnames:
        for row in original[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    written = saved[name][cell.coordinate].value
                    assert isinstance(written, str) and written.startswith("="), (
                        name,
                        cell.coordinate,
                        written,
                    )


# =============================================================================
# 2. Excel's own numbers, read back independently of the Checks sheet
# =============================================================================


@pytest.mark.parametrize(
    "case_name",
    ["golden_v2_1", "revenue_outgrows_expenses", "expenses_outgrow_revenue", "negative_expense_growth"],
)
def test_excels_operating_lines_equal_anchors_engine(case_name: str, recalculated) -> None:  # noqa: ANN001
    """Read Excel's operating cells directly and compare with the projection
    Anchor computed, without consulting a single Checks row."""

    case = CASES_BY_NAME[case_name]
    projection = analyze(case).operating_projection
    ws = load(recalculated[case_name], values=True)[OPERATING]
    start = row_of(ws, "Year number")
    fields = {
        "Gross potential rent": projection.gross_potential_rent_by_year,
        "Other income": projection.other_income_by_year,
        "Vacancy and credit loss": projection.vacancy_credit_loss_by_year,
        "Effective gross income": projection.effective_gross_income_by_year,
        "Property taxes": projection.property_taxes_by_year,
        "Insurance": projection.insurance_by_year,
        "Utilities": projection.utilities_by_year,
        "Repairs and maintenance": projection.repairs_maintenance_by_year,
        "Other operating expenses": projection.other_operating_expenses_by_year,
        "Management fee": projection.management_fee_by_year,
        "Total operating expenses": projection.total_operating_expenses_by_year,
        "Net operating income": projection.noi_by_year,
    }
    for label, values in fields.items():
        row = row_of(ws, label, start=start)
        for offset, expected in enumerate(values):
            actual = ws.cell(row, 3 + offset).value
            assert actual == pytest.approx(expected, rel=1e-12, abs=1e-6), (label, offset, actual, expected)


@pytest.mark.parametrize(
    "case_name", ["golden_v2_1", "revenue_outgrows_expenses", "expenses_outgrow_revenue"]
)
def test_excels_forward_year_noi_equals_anchors_exit_noi(case_name: str, recalculated) -> None:  # noqa: ANN001
    """The exit-only year, read from Excel's own cell. This is the number a
    blended growth rate would have got wrong."""

    case = CASES_BY_NAME[case_name]
    analysis = analyze(case)
    ws = load(recalculated[case_name], values=True)[OPERATING]
    noi_row = row_of(ws, "Net operating income", start=row_of(ws, "Year number"))
    forward = number_at(ws, noi_row, 3 + case.terms.hold_period)
    assert forward == pytest.approx(analysis.results.exit_noi, rel=1e-12, abs=1e-6)

    # And it is genuinely a further year of the schedule, not NOI_H rescaled.
    noi = analysis.operating_projection.noi_by_year
    if case_name != "golden_v2_1":
        blended = noi[-1] * (noi[-1] / noi[-2])
        assert abs(forward - blended) > 1.0, (forward, blended)


def test_excels_downstream_results_equal_anchors(recalculated) -> None:  # noqa: ANN001
    case = CASES_BY_NAME["golden_v2_1"]
    results = analyze(case).results
    wb = load(recalculated["golden_v2_1"], values=True)
    equity = wb["Equity Cash Flow"]
    assert equity.cell(row_of(equity, "Gross sale price  (exit NOI / exit cap rate)"), 3 + case.terms.hold_period).value == pytest.approx(
        results.exit_value, rel=1e-12, abs=1e-6
    )
    assert equity.cell(row_of(equity, "Equity multiple  (returned / invested)"), 3).value == pytest.approx(
        results.equity_multiple, rel=1e-12, abs=1e-9
    )
    assert equity.cell(row_of(equity, "Levered IRR"), 3).value == pytest.approx(
        results.levered_irr, rel=0, abs=1e-7
    )
    debt = wb["Debt Schedule"]
    assert debt.cell(row_of(debt, f"Loan payoff at sale (end of Year {case.terms.hold_period})"), 3).value == pytest.approx(
        results.remaining_loan_balance, rel=1e-12, abs=1e-6
    )


# =============================================================================
# 3. Mutation: each check must be able to fail
# =============================================================================


def _operating_cell(data: bytes, label: str, year: int) -> str:
    ws = load(data)[OPERATING]
    row = row_of(ws, label, start=row_of(ws, "Year number"))
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(2 + year)}{row}"


@pytest.mark.parametrize(
    ("label", "year", "broken", "guarded", "unrelated"),
    [
        # Vacancy taken on EGI instead of GPR: the convention says GPR only.
        (
            "Vacancy and credit loss",
            1,
            "=C{egi}*$C${pct}",
            "Vacancy and credit loss, Year 1",
            "Property taxes, Year 1",
        ),
        # Management fee grown like a fixed expense instead of scaled from EGI.
        (
            "Management fee",
            1,
            "=$C${fee_base}*C{expense_factor}",
            "Management fee, Year 1",
            "Gross potential rent, Year 1",
        ),
    ],
)
def test_breaking_an_operating_convention_fails_its_own_row(
    tmp_path: Path, label: str, year: int, broken: str, guarded: str, unrelated: str
) -> None:
    case = CASES_BY_NAME["golden_v2_1"]
    data = build(case)
    ws = load(data)[OPERATING]
    start = row_of(ws, "Year number")
    references = {
        "egi": row_of(ws, "Effective gross income", start=start),
        "pct": row_of(ws, "Vacancy and credit loss"),
        "fee_base": row_of(ws, "Management fee"),
        "expense_factor": row_of(ws, "Expense growth factor"),
    }
    reference = _operating_cell(data, label, year)
    mutated = replace_formula(data, OPERATING, reference, broken.format(**references))
    recalculated = _recalculate(tmp_path, {"mutant": mutated})["mutant"]
    statuses = _statuses(recalculated)
    assert not statuses[guarded].startswith("Pass"), (guarded, statuses[guarded])
    assert statuses[unrelated].startswith("Pass"), (unrelated, statuses[unrelated])


def test_approximating_exit_noi_with_a_blended_rate_fails_the_exit_row(tmp_path: Path) -> None:
    """The mutation the Projection Horizon convention exists to forbid: take
    Year H's NOI and grow it by the NOI series' own last-year rate instead of
    projecting the forward year. On a case where revenue and expense growth
    differ, the exit row must fail."""

    case = CASES_BY_NAME["revenue_outgrows_expenses"]
    data = build(case)
    hold = case.terms.hold_period
    ws = load(data)[OPERATING]
    noi_row = row_of(ws, "Net operating income", start=row_of(ws, "Year number"))
    from openpyxl.utils import get_column_letter

    last = get_column_letter(2 + hold)
    prior = get_column_letter(1 + hold)
    forward = f"{get_column_letter(3 + hold)}{noi_row}"
    mutated = replace_formula(
        data, OPERATING, forward, f"={last}{noi_row}*({last}{noi_row}/{prior}{noi_row})"
    )
    statuses = _statuses(_recalculate(tmp_path, {"blended": mutated})["blended"])
    assert not statuses["Exit NOI (Year H+1)"].startswith("Pass")
    assert not statuses["NOI identity (EGI - total operating expenses), Year 8 (exit year)"].startswith("Pass")
    assert statuses["Net operating income, Year 1"].startswith("Pass")


def test_a_deleted_formula_reads_missing_and_never_passes(tmp_path: Path) -> None:
    case = CASES_BY_NAME["golden_v2_1"]
    data = build(case)
    reference = _operating_cell(data, "Insurance", 2)
    mutated = blank_cell(data, OPERATING, reference)
    statuses = _statuses(_recalculate(tmp_path, {"deleted": mutated})["deleted"])
    assert statuses["Insurance, Year 2"] == "Missing" or not statuses["Insurance, Year 2"].startswith("Pass")
    assert statuses["Insurance, Year 1"].startswith("Pass")


def test_a_corrupted_formula_is_reported_as_an_excel_error_not_a_pass(tmp_path: Path) -> None:
    """An identity row's Anchor side is a live expression, so an error there
    must read as an error rather than escape into the counts."""

    case = CASES_BY_NAME["golden_v2_1"]
    data = build(case)
    reference = _operating_cell(data, "Utilities", 3)
    mutated = replace_formula(data, OPERATING, reference, "=1/0")
    recalculated = _recalculate(tmp_path, {"corrupt": mutated})["corrupt"]
    statuses = _statuses(recalculated)
    assert not statuses["Utilities, Year 3"].startswith("Pass")
    assert not statuses["Total operating expenses identity (six expense lines), Year 3"].startswith("Pass")
    assert statuses["Utilities, Year 1"].startswith("Pass")
    # The counts survive: an error is counted, never allowed to void them.
    block = status_block(load(recalculated, values=True))
    assert isinstance(block["Not passed"], int) and block["Not passed"] > 0


def test_editing_a_working_input_moves_its_dependents_and_nothing_else(tmp_path: Path) -> None:
    """A Working Input edit makes the workbook a modified case: every row says
    so, the dependent lines move, and an independent line does not."""

    case = CASES_BY_NAME["golden_v2_1"]
    data = build(case)
    ws = load(data)["Inputs"]
    header = row_of(ws, "Assumption")
    row = row_of(ws, "Insurance", start=header)
    original = number_at(ws, row, 2)
    mutated = set_number(data, "Inputs", f"C{row}", original * 2)
    recalculated = _recalculate(tmp_path, {"edited": mutated})["edited"]

    block = status_block(load(recalculated, values=True))
    assert block["Workbook modified since export"] == "Yes"
    assert block["Anchor reconciliation available"] == "No: Working Inputs modified"
    statuses = _statuses(recalculated)
    assert set(statuses.values()) == {"Not like-for-like"}

    # The model still moved: insurance and NOI changed, GPR did not.
    before = load(build(case))
    after = load(recalculated, values=True)[OPERATING]
    baseline = analyze(case).operating_projection
    start = row_of(after, "Year number")
    insurance_row = row_of(after, "Insurance", start=start)
    gpr_row = row_of(after, "Gross potential rent", start=start)
    noi_row = row_of(after, "Net operating income", start=start)
    assert after.cell(insurance_row, 3).value == pytest.approx(baseline.insurance_by_year[0] * 2)
    assert after.cell(gpr_row, 3).value == pytest.approx(baseline.gross_potential_rent_by_year[0])
    assert after.cell(noi_row, 3).value == pytest.approx(
        baseline.noi_by_year[0] - baseline.insurance_by_year[0]
    )
    assert before is not None


#: One perturbation per distinct class of Working Input, with the line it must
#: move and a line it must leave alone. ``(input label, new value, moved line,
#: moved year, unchanged line, unchanged year)``. Years are 1-based columns of
#: the operating statement; ``None`` for a moved line means a downstream
#: figure rather than an operating line.
_PERTURBATIONS = (
    ("Revenue growth", 0.12, "Gross potential rent", 3, "Property taxes", 3),
    ("Expense growth", 0.12, "Property taxes", 3, "Gross potential rent", 3),
    ("Vacancy and credit loss", 0.25, "Vacancy and credit loss", 2, "Property taxes", 2),
    ("Management fee", 0.20, "Management fee", 2, "Gross potential rent", 2),
    ("Property taxes", 500_000.0, "Property taxes", 1, "Insurance", 1),
)


def test_each_class_of_working_input_moves_its_own_dependents(tmp_path: Path) -> None:
    """Perturb each distinct class of operating input in turn, in one Excel
    session, and prove the lines it feeds move while an independent line does
    not."""

    case = CASES_BY_NAME["golden_v2_1"]
    data = build(case)
    inputs = load(data)["Inputs"]
    header = row_of(inputs, "Assumption")

    workbooks: dict[str, bytes] = {}
    for label, value, *_rest in _PERTURBATIONS:
        row = row_of(inputs, label, start=header)
        workbooks[label] = set_number(data, "Inputs", f"C{row}", value)
    recalculated = _recalculate(tmp_path, workbooks)

    baseline = analyze(case).operating_projection
    fields = {
        "Gross potential rent": baseline.gross_potential_rent_by_year,
        "Property taxes": baseline.property_taxes_by_year,
        "Insurance": baseline.insurance_by_year,
        "Vacancy and credit loss": baseline.vacancy_credit_loss_by_year,
        "Management fee": baseline.management_fee_by_year,
    }

    for label, _value, moved, moved_year, unchanged, unchanged_year in _PERTURBATIONS:
        ws = load(recalculated[label], values=True)[OPERATING]
        start = row_of(ws, "Year number")
        moved_value = ws.cell(row_of(ws, moved, start=start), 2 + moved_year).value
        unchanged_value = ws.cell(row_of(ws, unchanged, start=start), 2 + unchanged_year).value
        assert moved_value != pytest.approx(
            fields[moved][moved_year - 1], rel=1e-9
        ), (label, moved, moved_value)
        assert unchanged_value == pytest.approx(
            fields[unchanged][unchanged_year - 1], rel=1e-9
        ), (label, unchanged, unchanged_value)
        # NOI always moves: every one of these feeds it.
        noi = ws.cell(row_of(ws, "Net operating income", start=start), 2 + moved_year).value
        assert noi != pytest.approx(baseline.noi_by_year[moved_year - 1], rel=1e-9), label


def test_exit_cap_and_leverage_move_the_sale_and_the_debt_not_the_operating_lines(
    tmp_path: Path,
) -> None:
    """Two inputs that sit entirely below NOI: the operating statement must
    not move at all, and the figures they do feed must."""

    case = CASES_BY_NAME["golden_v2_1"]
    data = build(case)
    inputs = load(data)["Inputs"]
    header = row_of(inputs, "Assumption")
    workbooks = {
        "exit_cap": set_number(
            data, "Inputs", f"C{row_of(inputs, 'Exit cap rate', start=header)}", 0.09
        ),
        "leverage": set_number(
            data, "Inputs", f"C{row_of(inputs, 'Loan-to-value', start=header)}", 0.30
        ),
    }
    recalculated = _recalculate(tmp_path, workbooks)
    results = analyze(case).results
    baseline = analyze(case).operating_projection

    for name in workbooks:
        wb = load(recalculated[name], values=True)
        operating_ws = wb[OPERATING]
        start = row_of(operating_ws, "Year number")
        for label, values in (
            ("Gross potential rent", baseline.gross_potential_rent_by_year),
            ("Net operating income", baseline.noi_by_year),
        ):
            row = row_of(operating_ws, label, start=start)
            for offset, expected in enumerate(values):
                assert operating_ws.cell(row, 3 + offset).value == pytest.approx(
                    expected, rel=1e-12, abs=1e-6
                ), (name, label, offset)

    equity = load(recalculated["exit_cap"], values=True)["Equity Cash Flow"]
    gross = equity.cell(
        row_of(equity, "Gross sale price  (exit NOI / exit cap rate)"), 3 + case.terms.hold_period
    ).value
    assert gross != pytest.approx(results.exit_value, rel=1e-9)

    debt = load(recalculated["leverage"], values=True)["Debt Schedule"]
    loan = debt.cell(row_of(debt, "Loan amount"), 3).value
    assert loan == pytest.approx(case.terms.purchase_price * 0.30, rel=1e-12)
    assert loan != pytest.approx(results.loan_amount, rel=1e-9)


def test_one_business_plan_year_moves_only_that_year(tmp_path: Path) -> None:
    case = CASES_BY_NAME["business_plan_items"]
    data = build(case)
    ws = load(data)["Inputs"]
    working_row = row_of(ws, "Project capital - Working Input")
    from openpyxl.utils import get_column_letter

    # Year 2's Working Input, on the Business Plan block (first period column
    # is C).
    reference = f"{get_column_letter(4)}{working_row}"
    original = number_at(ws, working_row, 4)
    mutated = set_number(data, "Inputs", reference, original + 500_000.0)
    recalculated = _recalculate(tmp_path, {"bp": mutated})["bp"]

    operating_ws = load(recalculated, values=True)[OPERATING]
    start = row_of(operating_ws, "Year number")
    project_row = row_of(operating_ws, "Business Plan project capital", start=start)
    results = analyze(case).results
    assert operating_ws.cell(project_row, 4).value == pytest.approx(
        -(original + 500_000.0), rel=1e-12
    )
    assert operating_ws.cell(project_row, 3).value == pytest.approx(
        -results.project_capital_by_year[0], rel=1e-12
    )
    # NOI is above the line and must not move at all.
    noi_row = row_of(operating_ws, "Net operating income", start=start)
    for offset, expected in enumerate(results.noi_by_year):
        assert operating_ws.cell(noi_row, 3 + offset).value == pytest.approx(
            expected, rel=1e-12, abs=1e-6
        ), offset


def test_restoring_every_perturbed_value_returns_every_check_to_passing(tmp_path: Path) -> None:
    """The round trip: perturb, then set the value back, and the workbook is a
    like-for-like reconciliation again."""

    case = CASES_BY_NAME["golden_v2_1"]
    data = build(case)
    ws = load(data)["Inputs"]
    header = row_of(ws, "Assumption")
    row = row_of(ws, "Revenue growth", start=header)
    original = number_at(ws, row, 2)

    perturbed = set_number(data, "Inputs", f"C{row}", 0.12)
    restored = set_number(perturbed, "Inputs", f"C{row}", original)
    recalculated = _recalculate(tmp_path, {"perturbed": perturbed, "restored": restored})

    perturbed_block = status_block(load(recalculated["perturbed"], values=True))
    assert perturbed_block["Workbook modified since export"] == "Yes"

    restored_block = status_block(load(recalculated["restored"], values=True))
    assert restored_block["Workbook modified since export"] == "No"
    assert restored_block["Anchor reconciliation available"] == "Yes"
    assert restored_block["Not passed"] == 0
    statuses = {status for status in _statuses(recalculated["restored"]).values()}
    assert statuses == {"Pass"} or all(status.startswith("Pass") for status in statuses)


def test_a_perturbed_growth_rate_moves_every_later_year_but_not_year_1(tmp_path: Path) -> None:
    case = CASES_BY_NAME["golden_v2_1"]
    data = build(case)
    ws = load(data)["Inputs"]
    row = row_of(ws, "Revenue growth", start=row_of(ws, "Assumption"))
    mutated = set_number(data, "Inputs", f"C{row}", 0.10)
    after = load(_recalculate(tmp_path, {"growth": mutated})["growth"], values=True)[OPERATING]
    baseline = analyze(case).operating_projection
    start = row_of(after, "Year number")
    gpr_row = row_of(after, "Gross potential rent", start=start)
    assert after.cell(gpr_row, 3).value == pytest.approx(baseline.gross_potential_rent_by_year[0])
    assert after.cell(gpr_row, 4).value == pytest.approx(
        baseline.gross_potential_rent_by_year[0] * 1.10
    )
    assert after.cell(gpr_row, 4).value != pytest.approx(baseline.gross_potential_rent_by_year[1])


def test_tampering_with_an_original_export_value_is_reported(tmp_path: Path) -> None:
    case = CASES_BY_NAME["golden_v2_1"]
    data = build(case)
    ws = load(data)["Inputs"]
    row = row_of(ws, "Utilities", start=row_of(ws, "Assumption"))
    mutated = set_number(data, "Inputs", f"B{row}", 999_999.0)
    block = status_block(load(_recalculate(tmp_path, {"tampered": mutated})["tampered"], values=True))
    assert block["Original inputs unchanged"] == "No"
    assert str(block["Anchor reconciliation available"]).startswith("No")


def test_an_unrecalculated_workbook_shows_no_pass_anywhere() -> None:
    """No Excel involved: the file as written must already be honest."""

    statuses = set(_statuses(build(CASES_BY_NAME["golden_v2_1"])).values())
    assert statuses == {"Not recalculated"}


def test_the_recalculation_canary_cannot_be_faked() -> None:
    data = build(CASES_BY_NAME["golden_v2_1"])
    ws = load(data)[CHECKS]
    row = row_of(ws, "Excel formulas recalculated")
    assert read_formula(data, CHECKS, f"B{row}") == '"Yes"'
    assert load(data, values=True)[CHECKS].cell(row, 2).value == "Not recalculated"


def test_quick_and_detailed_agree_in_excel_on_the_convergence_case(tmp_path: Path) -> None:
    """The permanent convergence case, proved in the spreadsheet: the Detailed
    workbook's downstream results equal the Quick workbook's when their NOI
    schedules are the same series."""

    from anchor.contracts import AcquisitionInputs
    from excel_export_golden_cases import GoldenCase
    from excel_export_golden_cases import build as build_quick

    case = CASES_BY_NAME["golden_v2_1"]
    quick_case = GoldenCase(
        "convergence",
        AcquisitionInputs(
            purchase_price=case.terms.purchase_price,
            current_noi=600_000.0,
            occupancy=0.95,
            noi_growth=0.03,
            hold_period=case.terms.hold_period,
            exit_cap_rate=case.terms.exit_cap_rate,
            ltv=case.terms.ltv,
            interest_rate=case.terms.interest_rate,
            amortization=case.terms.amortization,
            acquisition_cost_pct=case.terms.acquisition_cost_pct,
            financing_fee_pct=case.terms.financing_fee_pct,
            disposition_cost_pct=case.terms.disposition_cost_pct,
            annual_capex_reserve=case.terms.annual_capex_reserve,
            io_period=case.terms.io_period,
        ),
    )
    recalculated = _recalculate(
        tmp_path, {"detailed": build(case), "quick": build_quick(quick_case)}
    )
    detailed = load(recalculated["detailed"], values=True)["Equity Cash Flow"]
    quick = load(recalculated["quick"], values=True)["Equity Cash Flow"]
    for label, tolerance in (
        ("Equity multiple  (returned / invested)", 1e-9),
        ("Levered IRR", 1e-7),
        ("Unlevered IRR", 1e-7),
        ("Total profit", 1e-6),
    ):
        detailed_value = detailed.cell(row_of(detailed, label), 3).value
        quick_value = quick.cell(row_of(quick, label), 3).value
        assert detailed_value == pytest.approx(quick_value, rel=0, abs=tolerance), label


def test_a_hostile_deal_name_stays_text_after_excel_opens_it(tmp_path: Path) -> None:
    hostile = '=cmd|\' /C calc\'!A0'
    source = dataclasses.replace(source_for(CASES_BY_NAME["golden_v2_1"]), deal_name=hostile)
    from anchor.exports.excel import build_detailed_audit_workbook

    recalculated = _recalculate(tmp_path, {"hostile": build_detailed_audit_workbook(source)})["hostile"]
    ws = load(recalculated, values=True)["Summary"]
    assert ws["A2"].value == hostile


def test_a_zero_value_case_recalculates_without_dividing_by_zero(tmp_path: Path) -> None:
    case = dataclasses.replace(
        CASES_BY_NAME["golden_v2_1"],
        name="all_zero_lines",
        terms=terms(ltv=0.0, io_period=0, annual_capex_reserve=0.0),
        operating=operating(
            other_income=0.0,
            property_taxes=0.0,
            insurance=0.0,
            utilities=0.0,
            repairs_maintenance=0.0,
            other_operating_expenses=0.0,
            management_fee_pct=0.0,
            vacancy_credit_loss_pct=0.0,
            revenue_growth=0.0,
            expense_growth=0.0,
        ),
    )
    data = _recalculate(tmp_path, {"zeros": build(case)})["zeros"]
    statuses = _statuses(data)
    assert {status for status in statuses.values() if not status.startswith("Pass")} == set()
    wb = load(data, values=True)
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for cell in row:
                assert not (isinstance(cell.value, str) and cell.value.startswith("#")), (
                    name,
                    cell.coordinate,
                    cell.value,
                )
