"""Excel Export 1 -- the Quick Underwrite formula-audit workbook: structure,
formula contract, security, eligibility, persistence read and API.

These run in the default suite and need no spreadsheet engine. What Excel
actually *calculates* from these formulas is proved separately, by native
recalculation, in ``tests/test_excel_export_1_native_recalc.py``. Nothing here
treats a cached value as evidence: the generated workbook deliberately carries
none, which is itself asserted below.
"""

from __future__ import annotations

import dataclasses
import io
import re
import sqlite3
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from openpyxl.worksheet.worksheet import Worksheet

from _d6_5_fixtures import create as create_mode_deal
from anchor import deals as deals_store
from anchor.api import app
from anchor.asset_types import AssetClassification, AssetType
from anchor.business_plan import BusinessPlan, CapitalItemCategory, CapitalPlanItem, OwnerExpenseCategory, OwnerExpenseItem
from anchor.contracts import OperatingMode, UnsupportedOperatingModeError
from anchor.deals import DealNotFoundError
from anchor.deals.fingerprint import fingerprint_quick_inputs
from anchor.deals.store import QuickAnalysisState, get_quick_analysis_provenance
from anchor.exports.excel import (
    MAX_EXPORT_HOLD_PERIOD,
    SHEET_ORDER,
    XLSX_MEDIA_TYPE,
    QuickAuditExportError,
    QuickAuditRefusalCode,
    build_quick_audit_workbook,
    content_disposition,
    quick_audit_filename,
    quick_audit_source,
    sanitize_deal_name,
)
from anchor.exports.excel import provenance as export_provenance
from excel_export_golden_cases import (
    CASES_BY_NAME,
    GENERATED_AT,
    GOLDEN_CASES,
    analyze,
    build,
    check_rows,
    load,
    row_of,
    source_for,
    status_block,
)

EXPECTED_SHEETS = [
    "Summary",
    "Inputs",
    "Operating Projection",
    "Debt Schedule",
    "Equity Cash Flow",
    "Anchor Results",
    "Checks",
    "Audit Metadata",
]
MODEL_SHEETS = ("Operating Projection", "Debt Schedule", "Equity Cash Flow")
GREEN, BLUE, GRAY = "FF00703C", "FF0000FF", "FF595959"


@pytest.fixture(scope="module")
def base_bytes() -> bytes:
    return build(CASES_BY_NAME["interest_only_with_fees"])


@pytest.fixture(scope="module")
def base(base_bytes: bytes) -> openpyxl.Workbook:
    return load(base_bytes)


@pytest.fixture(scope="module")
def plan_bytes() -> bytes:
    return build(CASES_BY_NAME["business_plan_items"])


def _formula_cells(wb: openpyxl.Workbook, sheet: str) -> list[Any]:
    return [cell for row in wb[sheet].iter_rows() for cell in row if cell.data_type == "f"]


def _all_formulas(wb: openpyxl.Workbook) -> list[tuple[str, str, str]]:
    return [
        (ws.title, cell.coordinate, str(cell.value))
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.data_type == "f"
    ]


def _defined_names(wb: openpyxl.Workbook) -> dict[str, str]:
    return {name: definition.attr_text for name, definition in wb.defined_names.items()}


# =============================================================================
# Workbook structure
# =============================================================================


def test_sheets_are_exactly_the_contract_in_order(base: openpyxl.Workbook) -> None:
    assert base.sheetnames == EXPECTED_SHEETS
    assert list(SHEET_ORDER) == EXPECTED_SHEETS
    active = base.active
    assert active is not None and active.title == "Summary"


@pytest.mark.parametrize("case", GOLDEN_CASES, ids=lambda case: case.name)
def test_every_golden_case_builds_the_same_sheet_set(case) -> None:  # noqa: ANN001
    assert load(build(case)).sheetnames == EXPECTED_SHEETS


def test_no_sheet_is_blank(base: openpyxl.Workbook) -> None:
    for ws in base.worksheets:
        assert ws.max_row > 5 and isinstance(ws["A1"].value, str), ws.title


def test_workbook_asks_for_a_full_automatic_recalculation_on_open(base_bytes: bytes) -> None:
    workbook_xml = zipfile.ZipFile(io.BytesIO(base_bytes)).read("xl/workbook.xml").decode()
    calc = re.search(r"<calcPr[^>]*/>", workbook_xml)
    assert calc is not None
    assert 'fullCalcOnLoad="1"' in calc.group(0)
    assert "calcMode" not in calc.group(0) or 'calcMode="auto"' in calc.group(0)


def test_package_has_no_macros_external_links_or_connections(base_bytes: bytes) -> None:
    names = zipfile.ZipFile(io.BytesIO(base_bytes)).namelist()
    forbidden = ("vbaProject", "externalLink", "connections", "queryTable", "activeX", "embeddings")
    assert not [name for name in names if any(token in name for token in forbidden)]
    content_types = zipfile.ZipFile(io.BytesIO(base_bytes)).read("[Content_Types].xml").decode()
    assert "macroEnabled" not in content_types


def test_formulas_use_no_volatile_external_or_dangerous_functions(base: openpyxl.Workbook) -> None:
    banned = re.compile(r"\b(INDIRECT|OFFSET|NOW|TODAY|RAND|RANDBETWEEN|CELL|INFO|HYPERLINK|WEBSERVICE|CALL|REGISTER\.ID)\s*\(", re.I)
    formulas = _all_formulas(base)
    assert len(formulas) > 500
    for sheet, ref, formula in formulas:
        assert not banned.search(formula), (sheet, ref, formula)
        assert "[" not in formula, (sheet, ref, formula)  # no external workbook reference


def test_every_sheet_is_protected_without_a_password_and_no_formula_is_hidden(base: openpyxl.Workbook) -> None:
    for ws in base.worksheets:
        assert ws.protection.sheet, ws.title
        assert not ws.protection.password, ws.title
        for cell in _formula_cells(base, ws.title):
            assert cell.protection.locked, (ws.title, cell.coordinate)
            assert not cell.protection.hidden, (ws.title, cell.coordinate)


def _working_input_cells(wb: openpyxl.Workbook) -> list[Any]:
    ws = wb["Inputs"]
    return [
        cell
        for row in ws.iter_rows()
        for cell in row
        if cell.font is not None and cell.font.color is not None and cell.font.color.rgb == BLUE
    ]


def test_only_working_inputs_are_unlocked_and_they_are_blue_numbers(base: openpyxl.Workbook) -> None:
    unlocked = [
        (ws.title, cell.coordinate)
        for ws in base.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None and not cell.protection.locked
    ]
    working = _working_input_cells(base)
    assert sorted(unlocked) == sorted(("Inputs", cell.coordinate) for cell in working)
    # 12 editable scalars + closing capital + 2 x 5 Business Plan years
    # (occupancy, hold period and post-hold capital are fixed at export).
    assert len(working) == 12 + 1 + 2 * 5
    for cell in working:
        assert isinstance(cell.value, (int, float)), cell.coordinate
    ws = base["Inputs"]
    header = row_of(ws, "Assumption")
    assert [ws.cell(header, col).value for col in range(1, 5)] == [
        "Assumption",
        "Original Export",
        "Working Input",
        "Units",
    ]


def test_working_inputs_start_equal_to_the_original_export(base: openpyxl.Workbook) -> None:
    ws = base["Inputs"]
    for cell in _working_input_cells(base):
        label = str(ws.cell(cell.row, 1).value)
        if label.endswith("- Working Input"):  # Business Plan by year: Original is the row above
            assert ws.cell(cell.row - 1, 1).value == label.replace("Working Input", "Original Export")
            assert ws.cell(cell.row - 1, cell.column).value == cell.value
        else:  # scalar block: Original Export is column B
            assert ws.cell(cell.row, 2).value == cell.value


def test_working_inputs_carry_anchors_input_domains_as_validation(base_bytes: bytes) -> None:
    xml = zipfile.ZipFile(io.BytesIO(base_bytes)).read("xl/worksheets/sheet2.xml").decode()
    validations = re.findall(r"<dataValidation [^>]*>", xml)
    assert len(validations) >= 12
    assert 'type="whole"' in xml  # amortization / IO years


def test_freeze_panes_and_readable_widths(base: openpyxl.Workbook) -> None:
    assert base["Inputs"].freeze_panes is not None
    assert base["Checks"].freeze_panes is not None
    for sheet in (*MODEL_SHEETS, "Anchor Results"):
        ws = base[sheet]
        assert ws.freeze_panes == "C1", sheet
        assert ws.column_dimensions["A"].width >= 55, sheet
        assert not ws.sheet_view.showGridLines, sheet


def test_labels_fit_their_column(base: openpyxl.Workbook) -> None:
    """A label longer than column A would be clipped by the units beside it."""

    for sheet in (*MODEL_SHEETS, "Anchor Results", "Inputs"):
        ws = base[sheet]
        width = ws.column_dimensions["A"].width
        for row in ws.iter_rows(min_row=3, max_col=2):
            label, neighbour = row[0].value, row[1].value
            if isinstance(label, str) and neighbour is not None:
                assert len(label) <= width, (sheet, row[0].coordinate, label)


def test_number_formats_follow_the_presentation_contract(base: openpyxl.Workbook) -> None:
    ws = base["Equity Cash Flow"]
    ecf = ws.cell(row_of(ws, "Total equity cash flow"), 3)
    assert ecf.number_format == '#,##0_);(#,##0);"-"_)'  # parentheses, dash for zero
    irr = ws.cell(row_of(ws, "Levered IRR"), 3)
    assert irr.number_format.startswith("0.00%")
    em = ws.cell(row_of(ws, "Equity multiple  (returned / invested)"), 3)
    assert '"x"' in em.number_format
    inputs = base["Inputs"]
    assert inputs.cell(row_of(inputs, "Loan-to-value"), 3).number_format.startswith("0.00%")
    assert inputs.cell(row_of(inputs, "Loan-to-value"), 3).value == pytest.approx(0.65)  # a decimal, not 65


def test_generated_timestamps_are_real_excel_dates(base: openpyxl.Workbook) -> None:
    summary = base["Summary"]
    generated = summary.cell(row_of(summary, "Generated (UTC)"), 2).value
    assert isinstance(generated, datetime)
    assert generated == GENERATED_AT.replace(tzinfo=None)
    audit = base["Audit Metadata"]
    iso = audit.cell(row_of(audit, "Generated (ISO 8601, with time zone)"), 2).value
    assert iso == "2026-09-19T15:30:00+00:00"


def test_anchor_results_are_numeric_constants_and_unavailable_is_never_zero() -> None:
    case = CASES_BY_NAME["full_leverage_no_equity"]
    results = analyze(case)
    wb = load(build(case))
    ws = wb["Anchor Results"]
    assert not _formula_cells(wb, "Anchor Results")
    assert ws.cell(row_of(ws, "Loan amount"), 3).value == results.loan_amount
    assert ws.cell(row_of(ws, "Equity multiple"), 3).value == "Unavailable"
    assert ws.cell(row_of(ws, "Levered IRR"), 3).value == "Unavailable"
    assert ws.cell(row_of(ws, "Levered IRR availability"), 3).value == "First nonzero cash flow not negative"
    assert ws.protection.sheet


def _same_to_16_digits(stored: object, saved: float) -> bool:
    """XlsxWriter stores numbers to 16 significant digits (documented in the
    workbook's Audit Metadata): the frozen value is the saved one to within one
    part in 10^15, never merely "close"."""

    return isinstance(stored, (int, float)) and abs(stored - saved) <= 1e-15 * max(abs(saved), 1e-300)


def test_anchor_results_hold_the_saved_values(base: openpyxl.Workbook) -> None:
    results = analyze(CASES_BY_NAME["interest_only_with_fees"])
    ws = base["Anchor Results"]
    assert _same_to_16_digits(ws.cell(row_of(ws, "Levered IRR"), 3).value, results.levered_irr or 0.0)
    assert _same_to_16_digits(ws.cell(row_of(ws, "Net sale proceeds"), 3).value, results.net_sale_proceeds)
    ecf = row_of(ws, "Equity cash flow")
    for t, saved in enumerate(results.levered_cash_flows):
        assert _same_to_16_digits(ws.cell(ecf, 3 + t).value, saved), t
    # The debt engine's annual balances end at the saved balance (the builder
    # checks bit equality before writing; see the inconsistency test).
    balance = row_of(ws, "Ending loan balance (see note)")
    assert _same_to_16_digits(ws.cell(balance, 3 + 5).value, results.remaining_loan_balance)
    audit = base["Audit Metadata"]
    assert "16 significant digits" in str(audit.cell(row_of(audit, "Stored precision"), 2).value)


def test_no_visible_internal_identifier_outside_audit_metadata(base: openpyxl.Workbook) -> None:
    deal_id = source_for(CASES_BY_NAME["interest_only_with_fees"]).deal_id
    for ws in base.worksheets:
        if ws.title == "Audit Metadata":
            continue
        for row in ws.iter_rows():
            for cell in row:
                assert cell.value != deal_id, (ws.title, cell.coordinate)
    audit = base["Audit Metadata"]
    assert audit.cell(row_of(audit, "Deal ID"), 2).value == deal_id


# =============================================================================
# Formula contract
# =============================================================================


def test_named_ranges_are_readable_and_point_at_working_inputs(base: openpyxl.Workbook) -> None:
    names = _defined_names(base)
    ws = base["Inputs"]
    expected = {
        "Purchase_Price": "Purchase price",
        "Current_NOI": "Current NOI (Year 1 NOI)",
        "NOI_Growth": "NOI growth",
        "Hold_Period": "Hold period",
        "Exit_Cap_Rate": "Exit cap rate",
        "Loan_To_Value": "Loan-to-value",
        "Interest_Rate": "Interest rate",
        "Amortization_Years": "Amortization",
        "IO_Period_Years": "Interest-only period",
        "Acquisition_Cost_Pct": "Acquisition costs",
        "Financing_Fee_Pct": "Financing fee",
        "Disposition_Cost_Pct": "Disposition costs",
        "CapEx_Reserve": "Annual CapEx reserve",
        "Closing_Project_Capital": "Closing project capital",
    }
    for name, label in expected.items():
        assert names[name] == f"'Inputs'!$C${row_of(ws, label)}", name


def _references(formula: str, names: dict[str, str]) -> set[tuple[str, str]]:
    """Every (sheet, cell) a formula reads, resolving defined names."""

    found: set[tuple[str, str]] = set()
    for sheet, a, b in re.findall(r"'([^']+)'!\$?([A-Z]+\$?\d+)(?::\$?([A-Z]+\$?\d+))?", formula):
        found.add((sheet, a.replace("$", "")))
        if b:
            found.add((sheet, b.replace("$", "")))
    for name, target in names.items():
        if re.search(rf"\b{name}\b", formula):
            sheet, cell = target.split("!")
            found.add((sheet.strip("'"), cell.replace("$", "")))
    return found


def test_model_sheets_read_only_working_inputs_never_originals_or_anchor(base: openpyxl.Workbook) -> None:
    names = _defined_names(base)
    working = {("Inputs", cell.coordinate) for cell in _working_input_cells(base)}
    inputs = base["Inputs"]
    # Structural fixed inputs (hold period) are formulas in the Working column.
    working.add(("Inputs", f"C{row_of(inputs, 'Hold period')}"))
    for sheet in MODEL_SHEETS:
        for cell in _formula_cells(base, sheet):
            references = _references(str(cell.value), names)
            for ref_sheet, ref in references:
                assert ref_sheet not in ("Anchor Results", "Checks", "Summary"), (sheet, cell.coordinate, cell.value)
                if ref_sheet == "Inputs":
                    assert (ref_sheet, ref) in working, (sheet, cell.coordinate, cell.value)


def test_anchor_results_are_read_only_by_checks_and_summary(base: openpyxl.Workbook) -> None:
    for sheet, ref, formula in _all_formulas(base):
        if "Anchor Results" in formula:
            assert sheet in ("Checks", "Summary"), (sheet, ref, formula)


def test_cell_colours_follow_one_rule(base: openpyxl.Workbook) -> None:
    """Green exactly when a formula reads another sheet (directly or through
    a name defined on another sheet); black otherwise; blue inputs; gray
    frozen Anchor values."""

    names = _defined_names(base)
    for ws in base.worksheets:
        for cell in _formula_cells(base, ws.title):
            refs = _references(str(cell.value), names)
            cross = any(ref_sheet != ws.title for ref_sheet, _ in refs)
            colour = cell.font.color.rgb if cell.font and cell.font.color else None
            if cross:
                assert colour == GREEN, (ws.title, cell.coordinate, cell.value, colour)
            else:
                assert colour != GREEN, (ws.title, cell.coordinate, cell.value, colour)
    anchor = base["Anchor Results"]
    value = anchor.cell(row_of(anchor, "Loan amount"), 3)
    assert value.font.color.rgb == GRAY


def test_representative_formulas_implement_the_quick_contract(base: openpyxl.Workbook) -> None:
    op = base["Operating Projection"]
    factor_row = row_of(op, "NOI growth factor  (1 + growth) ^ (year - 1)")
    noi_row = factor_row + 1
    year_row = row_of(op, "Year number")
    assert op.cell(factor_row, 4).value == f"=(1+$C$6)^(D${year_row}-1)"
    assert op.cell(noi_row, 4).value == f"=$C$5*D{factor_row}"
    assert op.cell(year_row, 8).value == 6  # forward column is Year H+1

    debt = base["Debt Schedule"]
    pmt = str(debt.cell(row_of(debt, "Amortizing payment  L x r / (1 - (1 + r)^-N)"), 3).value)
    assert "^(-" in pmt and "IF(" in pmt
    header = row_of(debt, "Month", start=row_of(debt, "Monthly schedule"))
    first = header + 1
    rate_row = row_of(debt, "Monthly rate  (annual rate / 12)")
    assert debt.cell(first, 6).value == f"=D{first}*$C${rate_row}"  # interest = balance x monthly rate
    assert debt.cell(first, 7).value == f"=E{first}-F{first}"  # principal = payment - interest
    assert str(debt.cell(first, 8).value).startswith(f"=IF(A{first}=$C$")  # zero at maturity
    assert debt.cell(header + 60, 1).value == 60 and debt.cell(header + 61, 1).value is None

    eq = base["Equity Cash Flow"]
    gross = eq.cell(row_of(eq, "Gross sale price  (exit NOI / exit cap rate)"), 8).value
    assert gross == "=$C$12/$C$8"
    assert eq.cell(row_of(eq, "Year 6 NOI (exit NOI)"), 3).value == f"='Operating Projection'!$H${noi_row}"
    ecf_row = row_of(eq, "Total equity cash flow")
    irr = str(eq.cell(row_of(eq, "Levered IRR"), 3).value)
    assert f"IRR($C${ecf_row}:$H${ecf_row}," in irr
    assert irr.startswith('=IF($C$') and '"Unavailable"' in irr
    em = eq.cell(row_of(eq, "Equity multiple  (returned / invested)"), 3).value
    invested = row_of(eq, "Total equity invested  (sum of negative periods)")
    returned = row_of(eq, "Total cash returned  (sum of positive periods)")
    assert em == f'=IF(C{invested}=0,"Unavailable",C{returned}/C{invested})'


def test_checks_cover_every_material_calculation(base: openpyxl.Workbook) -> None:
    ws = base["Checks"]
    metrics = {str(ws.cell(row, 1).value) for row in range(1, ws.max_row + 1)}
    required = [
        *(f"Net operating income, Year {y}" for y in range(1, 6)),
        *(f"CapEx reserve, Year {y}" for y in range(1, 6)),
        *(f"Scheduled debt service, Year {y}" for y in range(1, 6)),
        *(f"Ending loan balance, Year {y}" for y in range(1, 6)),
        *(f"Equity cash flow, Year {t}" for t in range(0, 6)),
        "Exit NOI (Year H+1)",
        "Exit value (gross sale price)",
        "Disposition costs",
        "Debt payoff at sale",
        "Net sale proceeds",
        "Total equity invested",
        "Total cash returned",
        "Total profit",
        "Equity multiple",
        "Levered IRR",
        "Levered IRR availability",
        "Unlevered IRR",
        "Hold period (years)",
        "Equity cash-flow periods",
        "Excel formulas recalculated",
        "Workbook modified since export",
        "Original inputs unchanged",
        "Anchor reconciliation available",
    ]
    missing = [metric for metric in required if metric not in metrics]
    assert missing == []
    header = row_of(ws, "Metric")
    assert [ws.cell(header, col).value for col in range(1, 7)] == [
        "Metric",
        "Anchor Result",
        "Excel Result",
        "Difference",
        "Tolerance",
        "Status",
    ]


def test_each_check_compares_a_frozen_anchor_value_with_a_model_cell(base: openpyxl.Workbook) -> None:
    ws = base["Checks"]
    header = row_of(ws, "Metric")
    compared = 0
    for row in range(header + 1, ws.max_row + 1):
        status = ws.cell(row, 6).value
        if not (isinstance(status, str) and status.startswith("=")):
            continue
        compared += 1
        anchor, excel = str(ws.cell(row, 2).value), str(ws.cell(row, 3).value)
        assert anchor.startswith("='Anchor Results'!"), (row, anchor)
        assert "Anchor Results" not in excel, (row, excel)
        assert "ISERROR" in status and '"Excel error"' in status
        assert '"Not like-for-like"' in status
        if "COUNT(" not in excel:
            assert excel.startswith("=IF(ISBLANK("), (row, excel)  # a deleted formula reads "Missing"
    assert compared > 80


def test_tolerances_are_explicit_and_narrow(base: openpyxl.Workbook) -> None:
    ws = base["Checks"]
    header = row_of(ws, "Metric")
    kinds: dict[str, Any] = {}
    for row in range(header + 1, ws.max_row + 1):
        metric = str(ws.cell(row, 1).value)
        if metric in ("Levered IRR", "Equity multiple", "Net sale proceeds", "Hold period (years)"):
            kinds[metric] = ws.cell(row, 5).value
    assert kinds["Levered IRR"] == 1e-7
    assert kinds["Equity multiple"] == 1e-10
    assert str(kinds["Net sale proceeds"]).startswith("=MAX(1e-06,1e-10*ABS(N(")
    assert kinds["Hold period (years)"] == "Exact"


def test_an_unrecalculated_workbook_never_passes_a_check(base_bytes: bytes) -> None:
    """Every formula's cached value is blank or 'Not recalculated': nothing is
    copied from Anchor, so only real recalculation can produce a Pass."""

    values = load(base_bytes, values=True)
    formulas = load(base_bytes)
    for ws in formulas.worksheets:
        for cell in _formula_cells(formulas, ws.title):
            cached = values[ws.title].cell(cell.row, cell.column).value
            assert cached in (None, "", "Not recalculated"), (ws.title, cell.coordinate, cached)
    rows = check_rows(values)
    assert rows and {status for *_, status in rows.values()} == {"Not recalculated"}
    block = status_block(values)
    assert block["Excel formulas recalculated"] == "Not recalculated"
    assert block["Anchor reconciliation available"] == "Not recalculated"


def test_workbook_is_deterministic_for_the_same_saved_deal() -> None:
    case = CASES_BY_NAME["business_plan_items"]
    assert build(case) == build(case)


def test_business_plan_items_are_listed_and_resolved_independently(plan_bytes: bytes) -> None:
    wb = load(plan_bytes)
    ws = wb["Inputs"]
    assert ws.cell(row_of(ws, "Capital c200"), 2).value == 200
    closing = row_of(ws, "Project capital from items")
    assert str(ws.cell(closing, 2).value).startswith("=SUMIFS(")
    assert "SUMPRODUCT(" in str(ws.cell(closing + 1, 3).value)
    rows = check_rows(load(plan_bytes, values=True))
    assert "Closing capital resolved from items" in rows
    assert "Owner expenses resolved from items, Year 6" in rows


# =============================================================================
# Column headers: one header, one column
# =============================================================================
#
# Headers were aligned to the data they label -- a value column's header right,
# a text column's header left -- so a right-aligned header and the left-aligned
# one beside it met at their shared cell boundary and read as a single run of
# words ("Working InputUnits"). A rule between them is too quiet to fix that on
# its own, so the separation is real space: a value header is centred, and a
# descriptive header that has a header to its left is indented away from it.
# The rule stays, marking the boundary inside that space.


HEADER_FILL = "FFDCE3EE"
HEADER_RULE = "FF8FA0BC"
#: One line of wrapped text, in points.
HEADER_LINE_HEIGHT = 13.5
#: One indent level is three spaces of the normal font (ECMA-376 alignment).
INDENT_CHARS = 3
#: XlsxWriter stores a column width plus Excel's own cell padding.
WIDTH_PADDING = 0.7109375
#: Clear space every adjacent header pair must keep, in character units.
MINIMUM_GAP = 2.0
#: Clear space the pairs that were reported as colliding must keep.
REPORTED_PAIR_GAP = 4.0


def _is_header(cell: Any) -> bool:
    return (
        cell.fill is not None
        and cell.fill.fgColor is not None
        and cell.fill.fgColor.rgb == HEADER_FILL
        and cell.font is not None
        and bool(cell.font.bold)
    )


def _header_bands(ws: Worksheet) -> dict[int, list[Any]]:
    """Every row of column headers on ``ws``, as row -> its header cells."""

    bands: dict[int, list[Any]] = {}
    for row in ws.iter_rows():
        headers: list[Any] = [cell for cell in row if _is_header(cell)]
        if headers:
            bands[int(headers[0].row)] = headers
    return bands


def _column_widths(data: bytes, sheet: str) -> dict[int, float]:
    """Nominal width per 1-based column, as the builder set it. Read from the
    sheet part because openpyxl registers only the first column of a span."""

    part = f"xl/worksheets/sheet{EXPECTED_SHEETS.index(sheet) + 1}.xml"
    xml = zipfile.ZipFile(io.BytesIO(data)).read(part).decode()
    widths: dict[int, float] = {}
    for attrs in re.findall(r"<col ([^>]*?)/?>", xml):
        found = dict(re.findall(r'(\w+)="([^"]*)"', attrs))
        if "width" not in found:
            continue
        for col in range(int(found["min"]), int(found["max"]) + 1):
            widths[col] = float(found["width"]) - WIDTH_PADDING
    return widths


def _greedy_lines(text: str, chars_per_line: int) -> int:
    """Lines needed to wrap ``text`` at word boundaries -- written here rather
    than imported, so the workbook's own arithmetic is not marking its own
    homework."""

    limit = max(1, chars_per_line)
    lines, current = 1, 0
    for word in text.split():
        length = len(word)
        while length > limit:
            lines += 1 if current == 0 else 2
            length -= limit
            current = 0
        if current == 0:
            current = length
        elif current + 1 + length <= limit:
            current += 1 + length
        else:
            lines += 1
            current = length
    return lines


class _Header:
    """One header cell, with the geometry that decides whether it reads as its
    own label: where its text sits inside its column."""

    def __init__(self, cell: Any, width: float) -> None:
        self.cell = cell
        self.column = int(cell.column)
        self.text = str(cell.value) if cell.value is not None else ""
        self.width = width
        self.align = cell.alignment.horizontal or "general"
        self.indent = int(cell.alignment.indent or 0)
        self.wrapped = bool(cell.alignment.wrap_text)

    @property
    def inset(self) -> float:
        """Character units between the cell's left edge and its text."""

        return INDENT_CHARS * self.indent

    def free_space(self) -> tuple[float, float]:
        """Clear character units to the left and right of the text, read
        conservatively: a wrapped header is assumed to fill its line."""

        if self.wrapped:
            return (0.0, 0.0) if self.align == "center" else (self.inset, 0.0)
        free = max(0.0, self.width - len(self.text))
        if self.align == "center":
            return free / 2, free / 2
        return self.inset, max(0.0, free - self.inset)


def _band(data: bytes, sheet: str, row: int) -> dict[str, _Header]:
    """One header band by header text."""

    ws = load(data)[sheet]
    widths = _column_widths(data, sheet)
    return {
        header.text: header
        for header in (_Header(cell, widths[int(cell.column)]) for cell in _header_bands(ws)[row])
    }


def _band_row(data: bytes, sheet: str, anchor_label: str) -> int:
    """The row of the band whose leading header is ``anchor_label`` -- located
    by its own text, never by a row number."""

    ws = load(data)[sheet]
    start = row_of(ws, "Monthly schedule") if anchor_label == "Month" else 1
    return row_of(ws, anchor_label, start=start)


def _every_band(data: bytes) -> list[tuple[str, int, list[Any]]]:
    wb = load(data)
    return [(ws.title, row, cells) for ws in wb.worksheets for row, cells in _header_bands(ws).items()]


def _adjacent_pairs(data: bytes) -> list[tuple[str, _Header, _Header, float]]:
    """Every neighbouring header pair in the workbook, with the clear space
    between their texts in character units."""

    wb = load(data)
    pairs: list[tuple[str, _Header, _Header, float]] = []
    for ws in wb.worksheets:
        widths = _column_widths(data, ws.title)
        for cells in _header_bands(ws).values():
            band = [_Header(cell, widths[int(cell.column)]) for cell in cells]
            for left, right in zip(band, band[1:]):
                if not left.text or not right.text:
                    continue  # a band written across a column with no header
                pairs.append((ws.title, left, right, left.free_space()[1] + right.free_space()[0]))
    return pairs


def test_every_column_header_is_ruled_off_from_the_one_beside_it(plan_bytes: bytes) -> None:
    bands = _every_band(plan_bytes)
    # Seven sheets carry column headers; Audit Metadata is label-and-value rows.
    assert {sheet for sheet, _, _ in bands} == set(EXPECTED_SHEETS) - {"Audit Metadata"}
    for sheet, row, cells in bands:
        columns = [cell.column for cell in cells]
        assert columns == list(range(columns[0], columns[0] + len(columns))), (sheet, row)
        for cell in cells:
            right = cell.border.right
            assert right is not None and right.style == "thin", (sheet, cell.coordinate)
            assert right.color is not None and right.color.rgb == HEADER_RULE, (sheet, cell.coordinate)
            assert cell.border.bottom is not None and cell.border.bottom.style == "thin", (sheet, cell.coordinate)


def test_headers_are_placed_so_their_text_leaves_the_shared_boundary(plan_bytes: bytes) -> None:
    """The mechanism, stated once: no header's text may sit against a boundary
    it shares with another header. A value header is centred (so it ends before
    its right edge and costs no width in the narrow value columns); a
    descriptive header with a header to its left is indented; only the leading
    label column, which nothing meets, keeps its text on the margin."""

    placed = 0
    for sheet, row, cells in _every_band(plan_bytes):
        widths = _column_widths(plan_bytes, sheet)
        band = [_Header(cell, widths[int(cell.column)]) for cell in cells]
        first = band[0].column
        for header in band:
            placed += 1
            if header.align == "center":
                assert header.indent == 0, (sheet, header.text)
            elif header.column == first:
                assert (header.align, header.indent) == ("left", 0), (sheet, row, header.text)
            else:
                assert header.align == "left", (sheet, header.text)
                assert header.indent >= 1, (sheet, header.text, "must be indented off the boundary")
            assert header.align in ("left", "center"), (sheet, header.text)
    assert placed > 50


def test_every_adjacent_header_pair_keeps_real_space_between_its_labels(plan_bytes: bytes) -> None:
    """A border is not separation. Every neighbouring pair keeps clear
    character space between the two texts, so they cannot read as one label
    even where the rule is missed."""

    pairs = _adjacent_pairs(plan_bytes)
    assert len(pairs) > 40
    tight = [(gap, sheet, left.text, right.text) for sheet, left, right, gap in pairs if gap < MINIMUM_GAP]
    assert tight == [], tight


#: The header pairs reported as running together, located by their own text:
#: a value header immediately followed by a descriptive one.
REPORTED_COLLISIONS = (
    ("Inputs", "Assumption", "Working Input", "Units"),
    ("Inputs", "Capital item", "Amount", "Category"),
    ("Inputs", "Owner-expense item", "Annual amount", "Last year (blank = through hold)"),
    ("Debt Schedule", "Month", "Hold year", "Phase"),
    ("Checks", "Metric", "Tolerance", "Status"),
    # Found on the same sweep: Summary's key-metric table has the same pair.
    ("Summary", "Metric", "Anchor (exported)", "Check"),
)


@pytest.mark.parametrize(
    ("sheet", "anchor_label", "value_header", "text_header"),
    REPORTED_COLLISIONS,
    ids=[f"{sheet}:{value}|{text}" for sheet, _, value, text in REPORTED_COLLISIONS],
)
def test_each_reported_header_pair_is_adjacent_and_pulled_apart(
    plan_bytes: bytes, sheet: str, anchor_label: str, value_header: str, text_header: str
) -> None:
    """The defect and its correction in one assertion: the two headers really
    are neighbours sharing a boundary, and each is now placed away from it --
    the value header centred, the descriptive header indented -- leaving space
    that reads unmistakably as two labels."""

    band = _band(plan_bytes, sheet, _band_row(plan_bytes, sheet, anchor_label))
    value, text = band[value_header], band[text_header]
    assert text.column == value.column + 1, (sheet, value_header, text_header)
    assert value.align == "center", (sheet, value_header)
    assert (text.align, text.indent) == ("left", 1), (sheet, text_header)
    gap = value.free_space()[1] + text.free_space()[0]
    assert gap >= REPORTED_PAIR_GAP, (sheet, value_header, text_header, gap)


def test_no_header_is_clipped_by_its_neighbour(plan_bytes: bytes) -> None:
    """A header either fits its own column, indent included, or wraps -- and a
    row holding a wrapped header is tall enough to show every line of it."""

    wrapped: list[tuple[str, str]] = []
    for sheet, row, cells in _every_band(plan_bytes):
        widths = _column_widths(plan_bytes, sheet)
        height = load(plan_bytes)[sheet].row_dimensions[row].height
        for cell in cells:
            header = _Header(cell, widths[int(cell.column)])
            if header.wrapped:
                wrapped.append((sheet, header.text))
                lines = _greedy_lines(header.text, int(header.width - header.inset))
                assert height is not None and height >= HEADER_LINE_HEIGHT * lines, (sheet, cell.coordinate, height)
                continue
            # Unwrapped, so it must fit: one character per unit of column width
            # is a generous reading of Excel's unit (a bold header character is
            # wider), so failing this would mean certain clipping.
            assert len(header.text) + header.inset <= header.width, (sheet, cell.coordinate, header.text)
    # Only genuinely long headers wrap: the rest keep the default row height.
    assert sorted(wrapped) == [
        ("Debt Schedule", "Beginning balance"),
        ("Debt Schedule", "Ending balance"),
        ("Inputs", "Last year (blank = through hold)"),
    ]


def test_every_audit_metadata_label_fits_its_column_or_wraps_into_a_tall_enough_row(
    plan_bytes: bytes,
) -> None:
    """Audit Metadata is label and value, not columns of headers. Its one
    over-long label -- the source commit, whose qualification is part of what
    it says -- wraps and its row grows, rather than being cut off by the value
    beside it or having the text shortened."""

    ws = load(plan_bytes)["Audit Metadata"]
    widths = _column_widths(plan_bytes, "Audit Metadata")
    label_width, value_width = widths[1], widths[2]
    wrapped: list[str] = []
    checked = 0
    for row in range(1, ws.max_row + 1):
        label, value = ws.cell(row, 1), ws.cell(row, 2)
        if not isinstance(label.value, str) or value.value is None:
            continue  # titles, notes and the navy section bars
        checked += 1
        text = label.value
        if not label.alignment.wrap_text:
            assert len(text) <= label_width, (label.coordinate, text, label_width)
            continue
        wrapped.append(text)
        lines = _greedy_lines(text, int(label_width))
        if value.alignment.wrap_text and isinstance(value.value, str):
            lines = max(lines, _greedy_lines(value.value, int(value_width)))
        height = ws.row_dimensions[row].height
        assert height is not None and height >= HEADER_LINE_HEIGHT * lines, (label.coordinate, height, lines)
    assert checked > 20
    assert wrapped == ["Source commit (checkout HEAD; uncommitted changes are not detected)"]
    # The qualification is preserved in full, not trimmed away.
    assert "uncommitted changes are not detected" in wrapped[0]


def test_only_a_sheet_title_a_wrapped_header_or_the_wrapped_audit_label_sets_a_row_height(
    plan_bytes: bytes,
) -> None:
    """The correction grew three rows and nothing else."""

    wb = load(plan_bytes)
    heights = {
        (ws.title, row): dimension.height
        for ws in wb.worksheets
        for row, dimension in ws.row_dimensions.items()
        if dimension.height is not None
    }
    titles = {(sheet, 1): 22.0 for sheet in EXPECTED_SHEETS}
    assert {key: height for key, height in heights.items() if key in titles} == titles
    grown = {key: height for key, height in heights.items() if key not in titles}
    audit = load(plan_bytes)["Audit Metadata"]
    source_commit = row_of(audit, "Source commit (checkout HEAD; uncommitted changes are not detected)")
    assert sorted(grown) == sorted(
        [("Debt Schedule", 38), ("Inputs", 45), ("Audit Metadata", source_commit)]
    )
    assert grown[("Debt Schedule", 38)] == grown[("Inputs", 45)] == 2 * HEADER_LINE_HEIGHT
    assert grown[("Audit Metadata", source_commit)] == 3 * HEADER_LINE_HEIGHT


def test_the_header_correction_widened_no_column_and_added_none(plan_bytes: bytes) -> None:
    """Presentation only: the column layout of every sheet is exactly what the
    accepted workbook had, so nothing was widened and no delimiter column was
    inserted between two headers."""

    label, units, period = 58, 21, 14
    expected = {
        "Summary": {1: 40, 2: 20, 3: 20, 4: 44},
        "Inputs": {1: label, 2: units, 3: units, 4: 30, 5: 36, **{col: period for col in range(6, 9)}},
        "Operating Projection": {1: label, 2: units, **{col: period for col in range(3, 10)}},
        "Debt Schedule": {1: label, 2: units, **{col: period for col in range(3, 9)}},
        "Equity Cash Flow": {1: label, 2: units, **{col: period for col in range(3, 10)}},
        "Anchor Results": {1: label, 2: units, **{col: period for col in range(3, 10)}},
        "Checks": {1: label, 2: 18, 3: 18, 4: 12, 5: 12, 6: 24, 7: 34, 8: 7},
        "Audit Metadata": {1: 36, 2: 110},
    }
    for sheet, widths in expected.items():
        stored = {col: round(width) for col, width in _column_widths(plan_bytes, sheet).items()}
        assert stored == widths, sheet


# =============================================================================
# Security: analyst-authored text is text
# =============================================================================


HOSTILE = ("=HYPERLINK(\"http://example.invalid\",\"x\")", "+SUM(1,2)", "-2+3", "@cmd|' /C calc'!A0")


def test_user_authored_text_is_written_as_strings_never_formulas() -> None:
    case = CASES_BY_NAME["amortizing"]
    plan = BusinessPlan(
        capital_items=(
            CapitalPlanItem(item_id="h1", description=HOSTILE[2], category=CapitalItemCategory.OTHER, month=3, amount=1.0),
        ),
        owner_expense_items=(
            OwnerExpenseItem(item_id="h2", description=HOSTILE[3], category=OwnerExpenseCategory.OTHER, annual_amount=2.0, first_year=1),
        ),
    )
    hostile_case = dataclasses.replace(case, business_plan=plan)
    source = dataclasses.replace(
        source_for(hostile_case),
        deal_name=HOSTILE[0],
        asset_subtype=HOSTILE[1],
    )
    data = build_quick_audit_workbook(source)
    wb = load(data)
    found = {
        str(cell.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.data_type == "s"
    }
    for text in HOSTILE:
        assert text in found, text
    for sheet, ref, formula in _all_formulas(wb):
        assert "HYPERLINK" not in formula and "cmd|" not in formula, (sheet, ref)


def test_no_local_path_or_environment_value_is_written(base_bytes: bytes) -> None:
    text = b"".join(
        zipfile.ZipFile(io.BytesIO(base_bytes)).read(name)
        for name in zipfile.ZipFile(io.BytesIO(base_bytes)).namelist()
    ).decode("utf-8", "replace")
    for marker in ("C:\\", "/Users/", "anchor.db", "ANCHOR_DB_PATH", "OPENAI", "AppData"):
        assert marker not in text, marker


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("Harbor Point", "Harbor Point"),
        ("../../etc/passwd", "-..-etc-passwd"),  # no separator survives
        ("a\\b/c:d*e?f\"g<h>i|j", "a-b-c-d-e-f-g-h-i-j"),
        ("  .hidden.  ", "hidden"),
        ("line\r\nbreak\x00nul", "linebreaknul"),
        ("", "Untitled Deal"),
        ("...", "Untitled Deal"),
        ("x" * 300, "x" * 100),
        ("Café Résidence", "Café Résidence"),
    ],
)
def test_filenames_are_sanitized_independently(name: str, expected: str) -> None:
    assert sanitize_deal_name(name) == expected
    filename = quick_audit_filename(name)
    assert filename == f"{expected} - Quick Underwrite Audit.xlsx"
    assert not re.search(r'[\\/:*?"<>|\x00-\x1f]', filename)


def test_content_disposition_is_an_attachment_with_safe_fallback() -> None:
    header = content_disposition(quick_audit_filename("Café \"Plaza\""))
    assert header.startswith("attachment; ")
    assert 'filename="Caf_ -Plaza- - Quick Underwrite Audit.xlsx"' in header
    assert "filename*=UTF-8''Caf%C3%A9%20-Plaza-%20-%20Quick%20Underwrite%20Audit.xlsx" in header
    assert "\r" not in header and "\n" not in header


def test_source_commit_reports_only_a_hash(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv(export_provenance.SOURCE_COMMIT_VARIABLE, "A" * 40)
    assert export_provenance.source_commit() == "a" * 40
    monkeypatch.setenv(export_provenance.SOURCE_COMMIT_VARIABLE, "not-a-hash; rm -rf /")
    commit = export_provenance.source_commit()
    assert commit is None or re.fullmatch(r"[0-9a-f]{40}", commit)
    assert export_provenance.anchor_version()


# =============================================================================
# Eligibility
# =============================================================================


def _provenance(case_name: str, state: QuickAnalysisState, **deal_overrides: Any):  # noqa: ANN202
    from anchor.deals.store import QuickAnalysisProvenance

    case = CASES_BY_NAME[case_name]
    results = analyze(case)
    deal = deals_store.Deal(
        id="deal-1",
        name="Eligibility",
        operating_mode=OperatingMode.QUICK,
        inputs=case.inputs,
        terms=None,
        detailed_operating_inputs=None,
        business_plan=case.business_plan,
        deal_context=None,
        analysis_snapshot=results if state is QuickAnalysisState.CURRENT else None,
        ai_snapshot=None,
        created_at=GENERATED_AT,
        updated_at=GENERATED_AT,
        **deal_overrides,
    )
    return QuickAnalysisProvenance(deal=deal, analysis_state=state, analysis_fingerprint="f" * 64)


def _source(provenance):  # noqa: ANN001, ANN202
    return quick_audit_source(provenance, generated_at=GENERATED_AT, anchor_version="test", source_commit=None)


@pytest.mark.parametrize(
    ("state", "code"),
    [
        (QuickAnalysisState.MISSING, QuickAuditRefusalCode.ANALYSIS_MISSING),
        (QuickAnalysisState.STALE, QuickAuditRefusalCode.ANALYSIS_STALE),
    ],
)
def test_missing_and_stale_analyses_are_refused_with_their_own_codes(state, code) -> None:  # noqa: ANN001
    with pytest.raises(QuickAuditExportError) as refusal:
        _source(_provenance("amortizing", state))
    assert refusal.value.code is code
    assert "save" in refusal.value.message.lower()


def test_a_current_analysis_becomes_a_source_with_its_classification() -> None:
    source = _source(
        _provenance("amortizing", QuickAnalysisState.CURRENT, asset_type=AssetType.OFFICE, asset_subtype="Suburban")
    )
    assert source.asset_type_label == "Office"
    assert source.asset_subtype == "Suburban"
    assert source.analysis_fingerprint == "f" * 64


def test_a_hold_period_beyond_the_export_limit_is_refused() -> None:
    case = CASES_BY_NAME["amortizing"]
    long_case = dataclasses.replace(case, inputs=dataclasses.replace(case.inputs, hold_period=MAX_EXPORT_HOLD_PERIOD + 1))
    from anchor.deals.store import QuickAnalysisProvenance

    results = analyze(long_case)
    provenance = _provenance("amortizing", QuickAnalysisState.CURRENT)
    deal = dataclasses.replace(provenance.deal, inputs=long_case.inputs, analysis_snapshot=results)
    with pytest.raises(QuickAuditExportError) as refusal:
        _source(QuickAnalysisProvenance(deal=deal, analysis_state=QuickAnalysisState.CURRENT, analysis_fingerprint="f"))
    assert refusal.value.code is QuickAuditRefusalCode.HOLD_PERIOD_EXCEEDS_EXPORT_LIMIT


def test_an_internally_inconsistent_saved_analysis_is_refused() -> None:
    source = source_for(CASES_BY_NAME["interest_only_with_fees"])
    tampered = dataclasses.replace(
        source, results=dataclasses.replace(source.results, remaining_loan_balance=source.results.remaining_loan_balance + 1.0)
    )
    with pytest.raises(QuickAuditExportError) as refusal:
        build_quick_audit_workbook(tampered)
    assert refusal.value.code is QuickAuditRefusalCode.ANALYSIS_INCONSISTENT


# =============================================================================
# Persistence read: provenance classification
# =============================================================================


def _save_quick(db: Path, case_name: str = "interest_only_with_fees", *, analysed: bool = True, name: str = "Harbor Point"):  # noqa: ANN202
    case = CASES_BY_NAME[case_name]
    deal = deals_store.create_deal(
        name,
        case.inputs,
        business_plan=case.business_plan,
        classification=AssetClassification(asset_type=AssetType.MULTIFAMILY, asset_subtype="Garden"),
        db_path=db,
    )
    if analysed:
        deal = deals_store.update_analysis_snapshot(
            deal.id,
            dataclasses.asdict(analyze(case)),
            financial_input_fingerprint=fingerprint_quick_inputs(case.inputs, business_plan=case.business_plan),
            db_path=db,
        )
    return deal


def test_provenance_distinguishes_current_missing_and_stale(tmp_path: Path) -> None:
    db = tmp_path / "anchor.db"
    unanalysed = _save_quick(db, analysed=False)
    assert get_quick_analysis_provenance(unanalysed.id, db_path=db).analysis_state is QuickAnalysisState.MISSING

    deal = _save_quick(db)
    current = get_quick_analysis_provenance(deal.id, db_path=db)
    assert current.analysis_state is QuickAnalysisState.CURRENT
    assert deal.inputs is not None
    assert current.analysis_fingerprint == fingerprint_quick_inputs(deal.inputs, business_plan=deal.business_plan)

    assert deal.inputs is not None
    deals_store.update_deal(
        deal.id, deal.name, dataclasses.replace(deal.inputs, noi_growth=0.05), business_plan=deal.business_plan, db_path=db
    )
    assert get_quick_analysis_provenance(deal.id, db_path=db).analysis_state is QuickAnalysisState.STALE


def test_provenance_refuses_other_modes_and_unknown_deals(tmp_path: Path) -> None:
    db = tmp_path / "anchor.db"
    for mode in ("detailed", "lease_level"):
        deal = create_mode_deal(mode, db, business_plan=BusinessPlan())
        with pytest.raises(UnsupportedOperatingModeError):
            get_quick_analysis_provenance(deal.id, db_path=db)
    with pytest.raises(DealNotFoundError):
        get_quick_analysis_provenance("no-such-deal", db_path=db)


# =============================================================================
# API
# =============================================================================


@pytest.fixture
def api_db(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    db = tmp_path / "anchor.db"
    monkeypatch.setenv("ANCHOR_DB_PATH", str(db))
    return db


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


def _url(deal_id: str) -> str:
    return f"/deals/{deal_id}/exports/quick-underwrite.xlsx"


def _dump(db: Path) -> list[str]:
    connection = sqlite3.connect(db)
    try:
        return list(connection.iterdump())
    finally:
        connection.close()


def test_export_downloads_the_saved_analysed_deal(api_db: Path, client: TestClient) -> None:
    deal = _save_quick(api_db, name="Harbor / Point: Phase 2")
    response = client.get(_url(deal.id))
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment; ")
    assert 'filename="Harbor - Point- Phase 2 - Quick Underwrite Audit.xlsx"' in disposition
    assert response.headers["cache-control"] == "no-store"
    wb = load(response.content)
    assert wb.sheetnames == EXPECTED_SHEETS
    summary = wb["Summary"]
    assert summary["A2"].value == "Harbor / Point: Phase 2"
    assert summary.cell(row_of(summary, "Asset Type"), 2).value == "Multifamily"
    audit = wb["Audit Metadata"]
    assert audit.cell(row_of(audit, "Analysis fingerprint"), 2).value == get_quick_analysis_provenance(
        deal.id, db_path=api_db
    ).analysis_fingerprint


def test_export_is_read_only(api_db: Path, client: TestClient) -> None:
    deal = _save_quick(api_db)
    before = _dump(api_db)
    for _ in range(2):
        assert client.get(_url(deal.id)).status_code == 200
    assert _dump(api_db) == before
    reread = deals_store.get_deal(deal.id, db_path=api_db)
    assert reread.updated_at == deal.updated_at
    assert reread.analysis_snapshot == deal.analysis_snapshot


def _refusal(response) -> tuple[int, str, str]:  # noqa: ANN001
    detail = response.json()["detail"]
    return response.status_code, detail["code"], detail["message"]


def test_missing_deal_is_a_typed_404(api_db: Path, client: TestClient) -> None:
    status, code, message = _refusal(client.get(_url("no-such-deal")))
    assert (status, code) == (404, "deal_not_found")
    assert "no-such-deal" not in message


@pytest.mark.parametrize(("mode", "label"), [("detailed", "Detailed Underwrite"), ("lease_level", "Lease-Level Underwrite")])
def test_other_modes_are_refused_as_unsupported(api_db: Path, client: TestClient, mode: str, label: str) -> None:
    deal = create_mode_deal(mode, api_db, business_plan=BusinessPlan())
    status, code, message = _refusal(client.get(_url(deal.id)))
    assert (status, code) == (422, "unsupported_operating_mode")
    assert label in message and "Quick Underwrite" in message


def test_missing_analysis_is_refused(api_db: Path, client: TestClient) -> None:
    deal = _save_quick(api_db, analysed=False)
    status, code, message = _refusal(client.get(_url(deal.id)))
    assert (status, code) == (409, "analysis_missing")
    assert "Analyze" in message


def test_stale_analysis_is_refused(api_db: Path, client: TestClient) -> None:
    deal = _save_quick(api_db)
    assert deal.inputs is not None
    deals_store.update_deal(
        deal.id, deal.name, dataclasses.replace(deal.inputs, exit_cap_rate=0.07), business_plan=deal.business_plan, db_path=api_db
    )
    status, code, _ = _refusal(client.get(_url(deal.id)))
    assert (status, code) == (409, "analysis_stale")


def test_generation_failure_is_typed_and_leaks_nothing(api_db: Path, client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    deal = _save_quick(api_db)

    def explode(_source: object) -> bytes:
        raise RuntimeError("C:\\Users\\secret\\anchor.db exploded")

    monkeypatch.setattr("anchor.api.build_quick_audit_workbook", explode)
    response = client.get(_url(deal.id))
    status, code, message = _refusal(response)
    assert (status, code) == (500, "export_generation_failed")
    assert "secret" not in response.text and "anchor.db" not in response.text and "RuntimeError" not in response.text
    assert message


def test_hold_period_limit_is_refused_through_the_api(api_db: Path, client: TestClient) -> None:
    case = CASES_BY_NAME["amortizing"]
    inputs = dataclasses.replace(case.inputs, hold_period=MAX_EXPORT_HOLD_PERIOD + 1)
    deal = deals_store.create_deal("Long hold", inputs, db_path=api_db)
    from anchor.analysis.business_plan_analysis import analyze_quick_acquisition_with_business_plan

    deals_store.update_analysis_snapshot(
        deal.id,
        dataclasses.asdict(analyze_quick_acquisition_with_business_plan(inputs, business_plan=BusinessPlan())),
        financial_input_fingerprint=fingerprint_quick_inputs(inputs),
        db_path=api_db,
    )
    status, code, _ = _refusal(client.get(_url(deal.id)))
    assert (status, code) == (422, "hold_period_exceeds_export_limit")


def test_the_web_client_may_read_the_download_filename(api_db: Path, client: TestClient) -> None:
    deal = _save_quick(api_db)
    response = client.get(_url(deal.id), headers={"Origin": "http://localhost:5173"})
    assert "content-disposition" in response.headers.get("access-control-expose-headers", "").lower()


def test_export_route_is_get_only(api_db: Path, client: TestClient) -> None:
    deal = _save_quick(api_db)
    for method in ("post", "put", "delete"):
        assert getattr(client, method)(_url(deal.id)).status_code == 405


def test_generation_uses_the_current_utc_time(api_db: Path, client: TestClient) -> None:
    deal = _save_quick(api_db)
    before = datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)
    wb = load(client.get(_url(deal.id)).content)
    summary = wb["Summary"]
    generated = summary.cell(row_of(summary, "Generated (UTC)"), 2).value
    assert isinstance(generated, datetime) and generated >= before.replace(second=0)
