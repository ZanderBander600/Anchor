"""Refinance V1 Stage 3 correction round -- Excel Exports 1-3 label the
acquisition-financing reference from the Deal's *true* owner.

A Unit of a visible Investment has no Base Capital Structure of its own: the
Deal route rightly refuses it (the Investment owns the structure). The first
Stage 3 cut read that route and swallowed the refusal, so such a Unit exported
without the acquisition-financing reference label even when its Investment
refinanced it. The owner is now resolved explicitly:

- a standalone Deal: its own Base structure;
- a Unit of a visible Investment: that Investment's Base structure;
- labeled exactly when an **executed** Base refinance applies to the Deal --
  one scoped to it, or one of the whole Investment -- never another Unit's;
- a structure that cannot be read is a typed ``capital_structure_unavailable``
  refusal, never an unlabelled workbook;
- no refinance: the accepted workbook, byte for byte.

Also here: ``unit_names_of`` names only a genuinely missing Unit "no longer
available"; corruption and unexpected errors propagate.
"""

from __future__ import annotations

import dataclasses
import io
import sqlite3
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import _refinance_v1_fixtures as rf  # type: ignore[import-not-found]
import _refinance_v1_stage_2_fixtures as fx  # type: ignore[import-not-found]
import _refinance_v1_stage_3_fixtures as s3  # type: ignore[import-not-found]
from _p7_6_fixtures import create_investment, detailed_deal, lease_level_deal  # type: ignore[import-not-found]
from anchor.analysis.business_plan_analysis import analyze_detailed_acquisition_with_business_plan
from anchor.capital_structure.events import AuthoredPositionRef, CapitalStructureWithEvents
from anchor.deals import refinance_presentation as presentation
from anchor.deals import store
from anchor.deals.contracts import DealNotFoundError
from anchor.deals.fingerprint import fingerprint_detailed_inputs
from anchor.deals.refinance_presentation import (
    AcquisitionReferenceUnavailableError,
    acquisition_reference_applies,
    unit_names_of,
)
from anchor.exports.excel import _workbook

ROUTES = {
    "quick": "quick-underwrite.xlsx",
    "detailed": "detailed-underwrite.xlsx",
    "lease_level": "lease-level.xlsx",
}
CORRUPTION = "UPDATE capital_events SET kind = 'recapitalization'"


@pytest.fixture
def db(tmp_path: Path) -> Path:
    return tmp_path / "anchor.db"


@pytest.fixture
def client(db: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    return s3.client(db, monkeypatch)


def _analysed_detailed(db: Path, deal: Any) -> Any:
    analysis = analyze_detailed_acquisition_with_business_plan(
        deal.terms, deal.detailed_operating_inputs, business_plan=deal.business_plan
    )
    return store.update_analysis_snapshot(
        deal.id,
        dataclasses.asdict(analysis),
        financial_input_fingerprint=fingerprint_detailed_inputs(
            deal.terms, deal.detailed_operating_inputs, business_plan=deal.business_plan
        ),
        db_path=db,
    )


def _deal(mode: str, db: Path, name: str) -> Any:
    """An export-eligible Deal of ``mode``: Quick and Detailed carry their
    saved analysis; Lease-Level analyses the saved state it reads."""

    if mode == "quick":
        deal = fx.base_deal(db, name=name)
        s3.analysed(db, deal)
        return deal
    if mode == "detailed":
        deal = detailed_deal(db, name=name)
        _analysed_detailed(db, deal)
        return deal
    return lease_level_deal(db, name=name)


def _export(client: TestClient, deal_id: str, mode: str) -> Any:
    return client.get(f"/deals/{deal_id}/exports/{ROUTES[mode]}")


def _labelled(content: bytes) -> bool:
    summary = load_workbook(io.BytesIO(content))["Summary"]
    text = [cell.value for row in summary.iter_rows() for cell in row if isinstance(cell.value, str)]
    return _workbook.REFINANCE_NOTICE in text


def _investment_scope_structure() -> CapitalStructureWithEvents:
    """A whole-Investment refinance over two unlevered Units (V1 defers an
    Investment refinance while Unit-scoped debt is outstanding)."""

    return CapitalStructureWithEvents(
        positions=(
            rf.closing_debt("sen-inv", amount=2_000_000.0, priority=2, rate=0.05, scope=rf.INVESTMENT_SCOPE),
            rf.replacement(
                "refi-inv",
                event_id="event-inv",
                scope=rf.INVESTMENT_SCOPE,
                priority=2,
                position_class=rf.PositionClass.MEZZANINE_DEBT,
            ),
        ),
        events=(
            rf.event(
                dscr=1.5,
                event_id="event-inv",
                replacement_id="refi-inv",
                scope=rf.INVESTMENT_SCOPE,
                retiring=(AuthoredPositionRef(position_id="sen-inv"),),
                label="Portfolio refinance",
            ),
        ),
    )


def _corrupt(db: Path) -> None:
    connection = sqlite3.connect(db)
    try:
        with connection:
            connection.execute(CORRUPTION)
    finally:
        connection.close()


# =============================================================================
# The seam: which Deal the reference applies to
# =============================================================================


def test_1_a_standalone_deal_without_a_refinance_is_not_labelled(db: Path) -> None:
    plain = fx.base_deal(db, name="Plain")
    store.set_deal_capital_structure(plain.id, fx.closing_mezz_only(plain.id), db_path=db)
    bare = fx.base_deal(db, name="Bare")

    assert acquisition_reference_applies(plain.id, db_path=db) is False
    assert acquisition_reference_applies(bare.id, db_path=db) is False


def test_1_a_standalone_deal_with_an_executed_refinance_is_labelled(db: Path) -> None:
    deal, _ = s3.refinance_deal(db)

    assert acquisition_reference_applies(deal.id, db_path=db) is True


def test_an_unexecuted_refinance_does_not_label(db: Path) -> None:
    """Only an executed refinance replaces the acquisition financing."""

    deal, _ = s3.refinance_deal(db, s3.f4_structure)

    assert acquisition_reference_applies(deal.id, db_path=db) is False


def test_2_3_a_unit_is_labelled_by_its_own_refinance_and_never_by_another_units(db: Path) -> None:
    a = fx.base_deal(db, name="Unit A")
    b = fx.base_deal(db, name="Unit B")
    investment = create_investment(db, a, b)
    store.set_base_capital_structure(investment.id, fx.evented(a.id, dscr=2.0), db_path=db)

    assert acquisition_reference_applies(a.id, db_path=db) is True
    assert acquisition_reference_applies(b.id, db_path=db) is False


def test_4_an_investment_scope_refinance_labels_every_unit(db: Path) -> None:
    a = fx.base_deal(db, name="Unit A", ltv=0.0)
    b = fx.base_deal(db, name="Unit B", ltv=0.0)
    investment = create_investment(db, a, b)
    store.set_base_capital_structure(investment.id, _investment_scope_structure(), db_path=db)

    assert acquisition_reference_applies(a.id, db_path=db) is True
    assert acquisition_reference_applies(b.id, db_path=db) is True


def test_a_strategy_replacement_is_never_consulted(db: Path) -> None:
    """The export is of the saved Base: a Strategy's refinance does not label it."""

    from anchor.analysis.strategy import InvestmentStrategyOverlay, StrategyDomain

    deal = fx.base_deal(db, name="Base without a refinance")
    investment_id, _ = store.set_deal_capital_structure(deal.id, fx.closing_mezz_only(deal.id), db_path=db)
    assert investment_id is not None
    store.create_strategy(
        investment_id,
        name="Refinance later",
        root_overlays=(InvestmentStrategyOverlay(domain=StrategyDomain.CAPITAL_STRUCTURE, content=s3.f5_structure(deal.id)),),
        db_path=db,
    )

    assert acquisition_reference_applies(deal.id, db_path=db) is False


def test_5_a_corrupt_owner_structure_fails_closed(db: Path) -> None:
    a = fx.base_deal(db, name="Unit A")
    b = fx.base_deal(db, name="Unit B")
    investment = create_investment(db, a, b)
    store.set_base_capital_structure(investment.id, fx.evented(a.id, dscr=2.0), db_path=db)
    standalone, _ = s3.refinance_deal(db, name="Standalone")
    _corrupt(db)

    for deal_id in (a.id, b.id, standalone.id):
        with pytest.raises(AcquisitionReferenceUnavailableError):
            acquisition_reference_applies(deal_id, db_path=db)


def test_a_missing_deal_is_its_own_error(db: Path) -> None:
    with pytest.raises(DealNotFoundError):
        acquisition_reference_applies("no-such-deal", db_path=db)


# =============================================================================
# The routes: Quick, Detailed and Lease-Level follow the same rule
# =============================================================================


@pytest.mark.parametrize("mode", sorted(ROUTES))
def test_6_every_mode_labels_unit_a_and_not_unit_b(mode: str, db: Path, client: TestClient) -> None:
    a = _deal(mode, db, "Unit A")
    b = _deal(mode, db, "Unit B")
    investment = create_investment(db, a, b)
    store.set_base_capital_structure(investment.id, fx.evented(a.id, dscr=2.0), db_path=db)

    labelled_a = _export(client, a.id, mode)
    labelled_b = _export(client, b.id, mode)

    assert labelled_a.status_code == 200, labelled_a.text
    assert labelled_b.status_code == 200, labelled_b.text
    assert _labelled(labelled_a.content) is True
    assert _labelled(labelled_b.content) is False


@pytest.mark.parametrize("mode", sorted(ROUTES))
def test_5_every_mode_refuses_a_corrupt_owner_structure(mode: str, db: Path, client: TestClient) -> None:
    a = _deal(mode, db, "Unit A")
    b = _deal(mode, db, "Unit B")
    investment = create_investment(db, a, b)
    store.set_base_capital_structure(investment.id, fx.evented(a.id, dscr=2.0), db_path=db)
    _corrupt(db)

    response = _export(client, b.id, mode)

    assert response.status_code == 409, response.text
    detail = response.json()["detail"]
    assert detail["code"] == "capital_structure_unavailable"
    assert a.id not in detail["message"] and investment.id not in detail["message"]


@pytest.mark.parametrize("mode", sorted(ROUTES))
def test_7_no_refinance_exports_the_accepted_workbook_byte_for_byte(
    mode: str, db: Path, client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The wrapper returns the very source it was given, so the workbook is the
    one the accepted builder always produced (the golden digest tests hold the
    builder itself)."""

    import anchor.api as api_module

    deal = _deal(mode, db, "No refinance")
    seen: list[tuple[Any, Any]] = []
    original = api_module._with_refinance_reference

    def spy(source: Any, refuse: Any) -> Any:
        result = original(source, refuse)
        seen.append((source, result))
        return result

    monkeypatch.setattr(api_module, "_with_refinance_reference", spy)
    response = _export(client, deal.id, mode)

    assert response.status_code == 200, response.text
    ((given, returned),) = seen
    assert returned is given
    assert getattr(returned, "refinance_configured") is False
    assert _labelled(response.content) is False


# =============================================================================
# Correction 3: only a missing Unit is "no longer available"
# =============================================================================


def test_a_missing_unit_keeps_its_fallback_wording(db: Path) -> None:
    deal = fx.base_deal(db, name="Harbor")

    names = unit_names_of((deal.id, "deleted-unit"), db_path=db)

    assert names == {deal.id: "Harbor", "deleted-unit": "the selected Unit (no longer available)"}


def test_persisted_data_corruption_is_not_worded_as_a_missing_unit(monkeypatch: pytest.MonkeyPatch) -> None:
    def corrupt(unit_id: str, *, db_path: Any = None) -> Any:
        raise store.PersistedDealDataError("stored Deal cannot be decoded")

    monkeypatch.setattr(presentation.store, "get_deal", corrupt)

    with pytest.raises(store.PersistedDealDataError):
        unit_names_of(("unit-a",))


def test_an_unexpected_error_is_not_swallowed(monkeypatch: pytest.MonkeyPatch) -> None:
    def broken(unit_id: str, *, db_path: Any = None) -> Any:
        raise RuntimeError("database is locked")

    monkeypatch.setattr(presentation.store, "get_deal", broken)

    with pytest.raises(RuntimeError, match="database is locked"):
        unit_names_of(("unit-a",))
