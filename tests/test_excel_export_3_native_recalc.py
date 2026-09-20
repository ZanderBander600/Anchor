"""Excel Export 3 -- what a real spreadsheet engine calculates.

Opt-in (``ANCHOR_EXCEL_NATIVE_RECALC=1``), because it drives desktop Excel.
Everything else about the workbook is proved without one in
``test_excel_export_3_lease_level_audit.py``; what needs Excel is the only
thing that cannot be faked: the *values* the formulas produce.

Four kinds of evidence:

1. **Reconciliation.** Every golden case recalculates and every check passes,
   with no Excel error and no ``####`` cell anywhere, and every formula still
   a formula afterwards.
2. **Independent read-back.** Excel's own suite-month, property-month, annual,
   exit, debt and return cells are compared with Anchor's engine in Python,
   without going through the Checks sheet -- so a Checks row that silently
   compared a cell with itself could not hide.
3. **Mutation.** Each invariant is broken in turn and the row that guards it
   must fail while an unrelated row still passes. A check that cannot fail is
   not a check.
4. **Perturbation.** A Working Input is edited and only the lines it feeds may
   move; restoring it returns every check to passing.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from excel_lease_level_export_golden_cases import (
    CASES_BY_NAME,
    GOLDEN_CASES,
    analyze,
    blank_cell,
    build,
    check_rows,
    load,
    read_formula,
    replace_formula,
    row_of,
    set_number,
    status_block,
)
from anchor.leasing.calendar import month_index
from excel_native_recalc import native_recalc_enabled, recalculate_with_excel

pytestmark = pytest.mark.skipif(
    not native_recalc_enabled(),
    reason="set ANCHOR_EXCEL_NATIVE_RECALC=1 to recalculate in desktop Excel",
)

MONTHLY_LEASING = "Monthly Leasing"
MONTHLY_RECOVERIES = "Monthly Recoveries"
MONTHLY_PROPERTY = "Monthly Property"
ANNUAL = "Annual Projection"
ROLLOVER = "Rollover"
INPUTS = "Inputs"

#: The case every mutation and perturbation is run against. It is chosen, and
#: built, so that **every line a mutation targets carries a real value**: free
#: rent on both branches, TI and LC, a renewal spread that makes the two
#: branches price differently, downtime, credit loss and partly recoverable
#: expenses. A probe whose target line is already zero would make its mutation
#: a no-op and its guard look untestable when it is not.
PROBE = "mutation_probe"


def _first_month_with_a_value(case, field: str) -> int:  # noqa: ANN001
    """The first canonical month where Anchor's own series is non-zero.

    Mutations are aimed with this rather than at a guessed month, so a
    mutation always changes a real number."""

    series = getattr(analyze(case).monthly_projection, field)
    for period, value in enumerate(series, start=1):
        if abs(value) > 1e-9:
            return period
    raise AssertionError(f"{case.name} has no non-zero {field} to mutate")


def _recalculate(tmp_path: Path, workbooks: dict[str, bytes]) -> dict[str, bytes]:
    """Recalculate each named workbook in one Excel session."""

    pairs = []
    for name, data in workbooks.items():
        source = tmp_path / f"{name}.xlsx"
        source.write_bytes(data)
        pairs.append((source, tmp_path / f"{name}.recalc.xlsx"))
    recalculate_with_excel(pairs, tmp_path)
    return {
        name: destination.read_bytes()
        for name, (_, destination) in zip(workbooks, pairs)
    }


def _statuses(data: bytes) -> dict[str, str]:
    return {metric: row[3] for metric, row in check_rows(load(data, values=True)).items()}


def _failing(data: bytes) -> dict[str, str]:
    return {
        metric: status
        for metric, status in _statuses(data).items()
        if not status.startswith("Pass")
    }


# =============================================================================
# 1. Every golden case reconciles in Excel
# =============================================================================


@pytest.fixture(scope="module")
def recalculated(tmp_path_factory: pytest.TempPathFactory) -> dict[str, bytes]:
    """Every golden case, recalculated once, in a single Excel session."""

    work = tmp_path_factory.mktemp("lease-level-native")
    return _recalculate(work, {case.name: build(case) for case in GOLDEN_CASES})


@pytest.mark.parametrize("case", GOLDEN_CASES, ids=lambda case: case.name)
def test_every_check_passes_after_a_real_recalculation(case, recalculated) -> None:  # noqa: ANN001
    data = recalculated[case.name]
    block = status_block(load(data, values=True))
    assert block["Excel formulas recalculated"] == "Yes"
    assert block["Anchor reconciliation available"] == "Yes"
    assert block["Workbook modified since export"] == "No"
    assert block["Original inputs unchanged"] == "Yes"

    statuses = _statuses(data)
    assert _failing(data) == {}
    assert block["Not passed"] == 0
    assert block["Passed"] == len(statuses)


@pytest.mark.parametrize("case", GOLDEN_CASES, ids=lambda case: case.name)
def test_no_cell_holds_an_excel_error(case, recalculated) -> None:  # noqa: ANN001
    wb = load(recalculated[case.name], values=True)
    errors = []
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    errors.append((name, cell.coordinate, cell.value))
    assert errors == []


@pytest.mark.parametrize("case", GOLDEN_CASES, ids=lambda case: case.name)
def test_no_column_is_too_narrow_for_its_widest_number(case, recalculated) -> None:
    """A ``####`` cell is a number Excel could not draw. It is never an error,
    so nothing else in this file would catch it."""

    wb = load(recalculated[case.name], values=True)
    for name in wb.sheetnames:
        ws = wb[name]
        for letter, dimension in ws.column_dimensions.items():
            width = dimension.width
            if width is None:
                continue
            for cell in ws[letter]:
                value = cell.value
                if not isinstance(value, (int, float)) or isinstance(value, bool):
                    continue
                # A currency cell renders as at most "(1,234,567,890)" plus a
                # little padding; anything that would need more room than the
                # column has is a ####.
                rendered = len(f"{abs(value):,.0f}") + (3 if value < 0 else 0)
                assert rendered <= width + 2, (name, cell.coordinate, value, width)


@pytest.mark.parametrize("case", GOLDEN_CASES, ids=lambda case: case.name)
def test_formulas_are_still_formulas_after_excel_saves(case, recalculated) -> None:  # noqa: ANN001
    original = load(build(case))
    saved = load(recalculated[case.name])
    for name in original.sheetnames:
        for row in original[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    written = saved[name][cell.coordinate].value
                    assert isinstance(written, str) and written.startswith("="), (
                        name,
                        cell.coordinate,
                        written,
                    )


def test_hostile_text_is_still_text_after_an_excel_round_trip(recalculated) -> None:  # noqa: ANN001
    wb = load(recalculated["hostile_text_values"], values=True)
    seen = []
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=cmd", "@SUM(A1:A9)", "+CMD")):
                    seen.append(cell.value)
    assert seen, "expected the hostile values to survive as text"


# =============================================================================
# 2. Excel's own numbers, read back independently of the Checks sheet
# =============================================================================


def _month_col(period: int) -> int:
    """The 1-based openpyxl column holding canonical month ``period``."""

    from anchor.exports.excel.lease_level_audit import _FIRST_MONTH_COL

    return _FIRST_MONTH_COL + period


@pytest.mark.parametrize(
    "case_name",
    ["multiple_suites", "initial_vacancy_market_lease_up", "free_rent_both_branches"],
)
def test_excels_suite_months_equal_anchors_engine(case_name: str, recalculated) -> None:  # noqa: ANN001
    """Every suite, every month, read straight out of Excel."""

    case = CASES_BY_NAME[case_name]
    monthly = analyze(case).monthly_projection
    ws = load(recalculated[case_name], values=True)[MONTHLY_LEASING]
    projections = {p.suite_id: p for p in monthly.operating_schedule.suite_projections}

    for suite in case.suites:
        label = suite.suite_label or suite.suite_id
        projection = projections[suite.suite_id]
        for field, line in (
            ("cash_base_rent", "Cash base rent"),
            ("contractual_base_rent", "Contractual base rent"),
            ("free_rent", "Free rent"),
            ("occupied_area_sf", "Occupied area"),
        ):
            row = row_of(ws, f"{label} -- {line}")
            for period, expected in enumerate(getattr(projection, field), start=1):
                actual = ws.cell(row, _month_col(period)).value
                assert actual == pytest.approx(expected, rel=1e-10, abs=1e-6), (
                    label, line, period, actual, expected,
                )


@pytest.mark.parametrize(
    "case_name",
    ["recovery_structures_all_three", "successor_modified_gross_with_stop", "partly_recoverable_expenses"],
)
def test_excels_monthly_recoveries_equal_anchors_engine(case_name: str, recalculated) -> None:  # noqa: ANN001
    case = CASES_BY_NAME[case_name]
    monthly = analyze(case).monthly_projection
    ws = load(recalculated[case_name], values=True)[MONTHLY_RECOVERIES]
    projections = {p.suite_id: p for p in monthly.recovery_schedule.suite_projections}
    for suite in case.suites:
        label = suite.suite_label or suite.suite_id
        row = row_of(ws, f"{label} -- Expense recoveries")
        for period, expected in enumerate(projections[suite.suite_id].expense_recovery, start=1):
            actual = ws.cell(row, _month_col(period)).value
            assert actual == pytest.approx(expected, rel=1e-10, abs=1e-6), (label, period)


@pytest.mark.parametrize(
    "case_name", ["multiple_suites", "credit_loss_and_management_fee", "growth_divergence"]
)
def test_excels_property_months_equal_anchors_engine(case_name: str, recalculated) -> None:  # noqa: ANN001
    case = CASES_BY_NAME[case_name]
    monthly = analyze(case).monthly_projection
    ws = load(recalculated[case_name], values=True)[MONTHLY_PROPERTY]
    for field, label in (
        ("cash_base_rent", "Cash base rent"),
        ("expense_recovery", "Expense recoveries"),
        ("other_income", "Other income"),
        ("credit_loss", "Credit loss"),
        ("effective_gross_income", "Effective gross income"),
        ("management_fee", "Management fee"),
        ("total_operating_expenses", "Total operating expenses"),
        ("noi", "Net operating income"),
        ("occupied_area_sf", "Occupied area"),
        ("vacant_area_sf", "Vacant area"),
    ):
        row = row_of(ws, label)
        for period, expected in enumerate(getattr(monthly, field), start=1):
            actual = ws.cell(row, _month_col(period)).value
            assert actual == pytest.approx(expected, rel=1e-10, abs=1e-6), (label, period)


@pytest.mark.parametrize("case_name", ["multiple_suites", "growth_divergence", "ten_year_hold_many_generations"])
def test_excels_annual_and_exit_equal_anchors_engine(case_name: str, recalculated) -> None:  # noqa: ANN001
    """The annual view and the forward exit window, read from Excel's own
    cells. The exit NOI is the number a blended growth rate would get wrong."""

    case = CASES_BY_NAME[case_name]
    annual = analyze(case).annual_projection
    ws = load(recalculated[case_name], values=True)[ANNUAL]
    hold = case.terms.hold_period
    for field, label in (
        ("noi_by_year", "Net operating income"),
        ("effective_gross_income_by_year", "Effective gross income"),
        ("expense_recovery_by_year", "Expense recoveries"),
        ("cash_base_rent_by_year", "Cash base rent"),
        ("management_fee_by_year", "Management fee"),
    ):
        row = row_of(ws, label)
        for offset, expected in enumerate(getattr(annual, field)):
            actual = ws.cell(row, 3 + offset).value
            assert actual == pytest.approx(expected, rel=1e-10, abs=1e-6), (label, offset)

    noi_row = row_of(ws, "Net operating income")
    exit_noi = ws.cell(noi_row, 3 + hold).value
    assert exit_noi == pytest.approx(annual.exit_noi, rel=1e-10, abs=1e-6)

    for field, label in (
        ("tenant_improvements_by_year", "Tenant improvements"),
        ("leasing_commissions_by_year", "Leasing commissions"),
    ):
        row = row_of(ws, label, start=row_of(ws, "Rent-roll leasing capital (below NOI)"))
        for offset, expected in enumerate(getattr(annual, field)):
            actual = ws.cell(row, 3 + offset).value
            assert actual == pytest.approx(expected, rel=1e-10, abs=1e-6), (label, offset)


@pytest.mark.parametrize("case_name", ["multiple_suites", "interest_only_debt", "all_cash"])
def test_excels_debt_and_returns_equal_anchors_engine(case_name: str, recalculated) -> None:  # noqa: ANN001
    """The shared downstream model, read back from Excel's own cells."""

    case = CASES_BY_NAME[case_name]
    results = analyze(case).results
    wb = load(recalculated[case_name], values=True)
    debt = wb["Debt Schedule"]
    equity = wb["Equity Cash Flow"]

    loan_row = row_of(debt, "Loan amount")
    assert debt.cell(loan_row, 3).value == pytest.approx(results.loan_amount, rel=1e-10, abs=1e-6)

    service_row = row_of(debt, "Scheduled debt service")
    for offset, expected in enumerate(results.annual_debt_service):
        actual = debt.cell(service_row, 3 + offset).value
        assert actual == pytest.approx(expected, rel=1e-10, abs=1e-6), offset

    for label, expected in (
        ("Total equity invested  (sum of negative periods)", results.total_equity_invested),
        ("Total cash returned  (sum of positive periods)", results.total_cash_returned),
        ("Total profit", results.total_profit),
    ):
        row = row_of(equity, label)
        assert equity.cell(row, 3).value == pytest.approx(expected, rel=1e-10, abs=1e-6), label


# =============================================================================
# 3. Mutation -- every guard must be able to fail
# =============================================================================


def _probe_cell(sheet: str, label: str, period: int, *, data: bytes | None = None) -> str:
    """The A1 reference of one month cell on a labelled row."""

    from openpyxl.utils import get_column_letter

    ws = load(data if data is not None else build(CASES_BY_NAME[PROBE]))[sheet]
    row = row_of(ws, label)
    return f"{get_column_letter(_month_col(period))}{row}"


@pytest.fixture(scope="module")
def probe_baseline(recalculated) -> bytes:  # noqa: ANN001
    return recalculated[PROBE]


def test_the_probe_case_passes_before_any_mutation(probe_baseline: bytes) -> None:
    """The control. Every mutation below is measured against this."""

    assert _failing(probe_baseline) == {}


def _mutate_and_recalculate(tmp_path: Path, mutations: dict[str, bytes]) -> dict[str, bytes]:
    return _recalculate(tmp_path, mutations)


@pytest.fixture(scope="module")
def mutants(tmp_path_factory: pytest.TempPathFactory) -> dict[str, bytes]:
    """Every mutation, built once and recalculated in one Excel session."""

    case = CASES_BY_NAME[PROBE]
    data = build(case)
    hold = case.terms.hold_period
    mid = 6 * hold  # a month comfortably inside the hold
    free_month = _first_month_with_a_value(case, "free_rent")
    ti_month = _first_month_with_a_value(case, "tenant_improvements")
    suite_label = case.suites[0].suite_label or case.suites[0].suite_id

    rent_cell = _probe_cell(MONTHLY_LEASING, f"{suite_label} -- Cash base rent", mid, data=data)
    free_cell = _probe_cell(MONTHLY_LEASING, f"{suite_label} -- Free rent", free_month, data=data)
    recovery_cell = _probe_cell(MONTHLY_RECOVERIES, f"{suite_label} -- Expense recoveries", mid, data=data)
    fee_cell = _probe_cell(MONTHLY_PROPERTY, "Management fee", mid, data=data)
    opex_cell = _probe_cell(MONTHLY_PROPERTY, "Property taxes", mid, data=data)
    ti_cell = _probe_cell(MONTHLY_PROPERTY, "Tenant improvements", ti_month, data=data)

    annual_ws = load(data)[ANNUAL]
    from openpyxl.utils import get_column_letter

    exit_cell = f"{get_column_letter(3 + hold)}{row_of(annual_ws, 'Net operating income')}"

    # The first rollover event's own probability mass.
    mass_cell = "H5"

    mutations: dict[str, bytes] = {
        # Rent: a suite's cash rent inflated by 1%.
        "rent": replace_formula(
            data, MONTHLY_LEASING, rent_cell, f"({read_formula(data, MONTHLY_LEASING, rent_cell)})*1.01"
        ),
        # Free rent: the concession silently dropped.
        "free_rent": replace_formula(data, MONTHLY_LEASING, free_cell, "0"),
        # Rollover: an event's probability mass forced to 1.
        "rollover_mass": replace_formula(data, ROLLOVER, mass_cell, "1"),
        # Recoveries: a suite's reimbursement halved.
        "recoveries": replace_formula(
            data, MONTHLY_RECOVERIES, recovery_cell,
            f"({read_formula(data, MONTHLY_RECOVERIES, recovery_cell)})/2",
        ),
        # Expenses: one fixed line grown at the wrong rate.
        "expenses": replace_formula(
            data, MONTHLY_PROPERTY, opex_cell, f"({read_formula(data, MONTHLY_PROPERTY, opex_cell)})*1.05"
        ),
        # The management fee taken on cash rent alone, excluding recoveries --
        # the classic circularity-avoidance mistake.
        "fee_basis": replace_formula(
            data, MONTHLY_PROPERTY, fee_cell,
            f"{_probe_cell(MONTHLY_PROPERTY, 'Cash base rent', mid, data=data)}*Management_Fee_Pct",
        ),
        # TI/LC: a leasing cheque dropped.
        "leasing_capital": replace_formula(data, MONTHLY_PROPERTY, ti_cell, "0"),
        # Exit NOI approximated by growing Year H instead of summing the
        # forward twelve months -- the one thing the horizon convention forbids.
        "exit_blended": replace_formula(
            data, ANNUAL, exit_cell,
            f"{get_column_letter(2 + hold)}{row_of(annual_ws, 'Net operating income')}*1.03",
        ),
        # A deleted formula must read Missing, never zero.
        "deleted": blank_cell(data, MONTHLY_PROPERTY, fee_cell),
        # An Excel error must stay visible as an Excel error.
        "errored": replace_formula(data, MONTHLY_PROPERTY, fee_cell, "1/0"),
    }
    work = tmp_path_factory.mktemp("lease-level-mutants")
    return _mutate_and_recalculate(work, mutations)


@pytest.mark.parametrize(
    ("mutant", "guarded"),
    [
        ("rent", "Cash base rent"),
        ("free_rent", "Free rent"),
        ("rollover_mass", "Cash base rent"),
        ("recoveries", "Expense recoveries"),
        ("expenses", "Property taxes"),
        ("fee_basis", "Management fee"),
        ("leasing_capital", "Tenant improvements"),
        ("exit_blended", "Exit NOI"),
    ],
)
def test_each_mutation_fails_the_row_that_guards_it(mutant: str, guarded: str, mutants) -> None:  # noqa: ANN001
    failing = _failing(mutants[mutant])
    assert failing, f"{mutant} changed a real number and must fail some row"
    assert any(guarded in metric for metric in failing), (mutant, guarded, sorted(failing)[:8])


@pytest.mark.parametrize(
    "mutant",
    ["rent", "free_rent", "rollover_mass", "recoveries", "expenses", "fee_basis", "leasing_capital"],
)
def test_an_unrelated_row_still_passes_under_each_mutation(mutant: str, mutants) -> None:  # noqa: ANN001
    """A mutation that failed *everything* would prove nothing about the row
    that is supposed to guard it."""

    statuses = _statuses(mutants[mutant])
    assert any(status.startswith("Pass") for status in statuses.values())
    assert statuses["Hold period (years)"].startswith("Pass")


def test_a_deleted_formula_reads_missing_never_zero(mutants) -> None:  # noqa: ANN001
    """A blank cell must read ``Missing`` in the Excel column -- never zero --
    and the row must fail. Reading it as zero would let a deleted formula
    reconcile against an Anchor value that happened to be zero."""

    rows = check_rows(load(mutants["deleted"], values=True))
    missing = {metric: row for metric, row in rows.items() if row[1] == "Missing"}
    assert missing, "a blank cell must read Missing"
    for metric, row in missing.items():
        assert row[1] != 0, metric
        assert not str(row[3]).startswith("Pass"), (metric, row[3])


def test_an_excel_error_stays_visible_as_an_excel_error(mutants) -> None:  # noqa: ANN001
    statuses = _statuses(mutants["errored"])
    assert any(status == "Excel error" for status in statuses.values())
    # The counts survive an error rather than being voided by it.
    block = status_block(load(mutants["errored"], values=True))
    assert isinstance(block["Passed"], (int, float))
    assert block["Passed"] > 0


# =============================================================================
# 4. Perturbation -- a Working Input moves what it feeds, and nothing else
# =============================================================================


def _input_cell(data: bytes, label: str) -> str:
    ws = load(data)[INPUTS]
    return f"C{row_of(ws, label, column=1)}"


@pytest.fixture(scope="module")
def perturbed(tmp_path_factory: pytest.TempPathFactory) -> dict[str, bytes]:
    case = CASES_BY_NAME[PROBE]
    data = build(case)
    edits = {
        "market_rent": (_input_cell(data, "Market rent"), 45.0),
        "expense_growth": (_input_cell(data, "Expense growth"), 0.09),
        "management_fee": (_input_cell(data, "Management fee"), 0.09),
        "renewal_probability": (_input_cell(data, "Renewal probability"), 0.1),
    }
    workbooks = {name: set_number(data, INPUTS, ref, value) for name, (ref, value) in edits.items()}
    # Restoring the edited value must return every check to passing.
    restore_ref, _ = edits["management_fee"]
    original = load(data)[INPUTS][restore_ref].value
    workbooks["restored"] = set_number(
        set_number(data, INPUTS, restore_ref, 0.09), INPUTS, restore_ref, float(original)
    )
    work = tmp_path_factory.mktemp("lease-level-perturbed")
    return _recalculate(work, workbooks)


@pytest.mark.parametrize(
    "edit", ["market_rent", "expense_growth", "management_fee", "renewal_probability"]
)
def test_editing_a_working_input_marks_the_workbook_modified(edit: str, perturbed) -> None:  # noqa: ANN001
    block = status_block(load(perturbed[edit], values=True))
    assert block["Workbook modified since export"] == "Yes"
    assert block["Working inputs that differ from the export"] == 1
    statuses = _statuses(perturbed[edit])
    assert all(status == "Not like-for-like" for status in statuses.values()), (
        edit, sorted(set(statuses.values())),
    )


def test_expense_growth_moves_expenses_and_noi_but_not_rent(perturbed) -> None:  # noqa: ANN001
    """The lines an input feeds must move, and an independent line must not."""

    edited = load(perturbed["expense_growth"], values=True)
    original = load(_baseline_values(), values=True)
    period = 24

    def value(wb, label: str) -> float:  # noqa: ANN001
        ws = wb[MONTHLY_PROPERTY]
        return ws.cell(row_of(ws, label), _month_col(period)).value

    assert value(edited, "Property taxes") != pytest.approx(value(original, "Property taxes"))
    assert value(edited, "Net operating income") != pytest.approx(
        value(original, "Net operating income")
    )
    # Revenue is untouched by an expense assumption.
    assert value(edited, "Cash base rent") == pytest.approx(value(original, "Cash base rent"))
    assert value(edited, "Other income") == pytest.approx(value(original, "Other income"))


def test_renewal_probability_moves_the_rollover_but_not_the_in_place_lease(perturbed) -> None:  # noqa: ANN001
    case = CASES_BY_NAME[PROBE]
    suite_label = case.suites[0].suite_label or case.suites[0].suite_id
    edited = load(perturbed["renewal_probability"], values=True)
    original = load(_baseline_values(), values=True)
    ws_edited, ws_original = edited[MONTHLY_LEASING], original[MONTHLY_LEASING]

    in_place_row = row_of(ws_original, "In-place contractual base rent")
    early = _month_col(3)
    assert ws_edited.cell(in_place_row, early).value == pytest.approx(
        ws_original.cell(in_place_row, early).value
    )

    # After the roll the probability split changes the expectation. Compared
    # over the whole post-roll window rather than one month: two branches can
    # coincide in any single month without the series being the same.
    total_row = row_of(ws_original, f"{suite_label} -- Cash base rent")
    months = 12 * case.terms.hold_period + 12
    moved = [
        period
        for period in range(1, months + 1)
        if ws_edited.cell(total_row, _month_col(period)).value
        != pytest.approx(ws_original.cell(total_row, _month_col(period)).value)
    ]
    assert moved, "changing the renewal probability must move the expected rent somewhere"
    # Nothing before the in-place lease expires may move: those months are
    # deterministic contractual history and are never probability-weighted.
    expiry = month_index(
        case.leases[0].lease_expiration_date,
        analysis_start=case.property_inputs.analysis_start_date,
    )
    assert min(moved) > expiry, (min(moved), expiry)


def test_restoring_a_working_input_returns_every_check_to_passing(perturbed) -> None:  # noqa: ANN001
    block = status_block(load(perturbed["restored"], values=True))
    assert block["Workbook modified since export"] == "No"
    assert _failing(perturbed["restored"]) == {}


_BASELINE: dict[str, bytes] = {}


def _baseline_values() -> bytes:
    """The unedited probe workbook, recalculated once and reused."""

    return _BASELINE["data"]


@pytest.fixture(scope="module", autouse=True)
def _prepare_baseline(recalculated) -> None:  # noqa: ANN001
    _BASELINE["data"] = recalculated[PROBE]
