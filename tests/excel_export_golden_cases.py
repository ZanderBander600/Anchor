"""Excel Export 1 test support: golden Quick Underwrite cases and helpers.

Each case is analysed by Anchor's own engine entry point -- exactly what a
saved analysis snapshot holds -- and its expected IRR statuses are asserted,
so a case cannot silently stop exercising the boundary it is named for.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from xml.sax.saxutils import escape

import openpyxl

from anchor.analysis.business_plan_analysis import analyze_quick_acquisition_with_business_plan
from anchor.business_plan import (
    BusinessPlan,
    CapitalItemCategory,
    CapitalPlanItem,
    OwnerExpenseCategory,
    OwnerExpenseItem,
)
from anchor.contracts import AcquisitionInputs
from anchor.deals.fingerprint import fingerprint_quick_inputs
from anchor.engine.contracts import AcquisitionResults, IrrStatus
from anchor.exports.excel import SHEET_ORDER, QuickAuditSource, build_quick_audit_workbook

GENERATED_AT = datetime(2026, 9, 19, 15, 30, 0, tzinfo=timezone.utc)

_BASE = dict(
    purchase_price=10_000_000.0,
    current_noi=600_000.0,
    occupancy=0.95,
    noi_growth=0.03,
    hold_period=5,
    exit_cap_rate=0.065,
    ltv=0.65,
    interest_rate=0.055,
    amortization=30,
)


def _inputs(**overrides: float | int) -> AcquisitionInputs:
    return AcquisitionInputs(**{**_BASE, **overrides})  # type: ignore[arg-type]


def _capital(item_id: str, month: int, amount: float) -> CapitalPlanItem:
    return CapitalPlanItem(
        item_id=item_id,
        description=f"Capital {item_id}",
        category=CapitalItemCategory.VALUE_ADD_RENOVATION,
        month=month,
        amount=amount,
    )


def _expense(item_id: str, amount: float, first: int, last: int | None) -> OwnerExpenseItem:
    return OwnerExpenseItem(
        item_id=item_id,
        description=f"Expense {item_id}",
        category=OwnerExpenseCategory.ASSET_MANAGEMENT,
        annual_amount=amount,
        first_year=first,
        last_year=last,
    )


@dataclass(frozen=True)
class GoldenCase:
    name: str
    inputs: AcquisitionInputs
    business_plan: BusinessPlan = field(default_factory=BusinessPlan)
    levered_status: IrrStatus = IrrStatus.DEFINED
    unlevered_status: IrrStatus = IrrStatus.DEFINED
    #: Check rows Excel is *documented* not to reproduce (Anchor-only numerical
    #: outcomes). Every other row must pass.
    expected_failures: frozenset[str] = frozenset()


GOLDEN_CASES: tuple[GoldenCase, ...] = (
    GoldenCase("amortizing", _inputs()),
    GoldenCase(
        "interest_only_with_fees",
        _inputs(
            io_period=2,
            acquisition_cost_pct=0.02,
            financing_fee_pct=0.01,
            disposition_cost_pct=0.02,
            annual_capex_reserve=50_000.0,
        ),
    ),
    GoldenCase("zero_growth", _inputs(noi_growth=0.0, hold_period=7, ltv=0.6, amortization=25)),
    GoldenCase("negative_growth", _inputs(noi_growth=-0.02, exit_cap_rate=0.07)),
    GoldenCase("strong_growth", _inputs(noi_growth=0.045, hold_period=10, interest_rate=0.0625)),
    GoldenCase("interest_only_beyond_hold", _inputs(io_period=10, hold_period=5)),
    GoldenCase("amortized_before_sale", _inputs(amortization=3, ltv=0.5)),
    GoldenCase("all_cash", _inputs(ltv=0.0, financing_fee_pct=0.01)),
    GoldenCase("zero_interest_rate", _inputs(interest_rate=0.0, financing_fee_pct=0.005)),
    GoldenCase(
        "full_leverage_no_equity",
        _inputs(ltv=1.0, current_noi=900_000.0),
        levered_status=IrrStatus.FIRST_NONZERO_NOT_NEGATIVE,
    ),
    GoldenCase(
        "business_plan_items",
        _inputs(hold_period=6),
        BusinessPlan(
            capital_items=(
                _capital("c0", 0, 250_000.0),
                _capital("c6", 6, 1_500_000.0),
                _capital("c14", 14, 125_000.0),
                _capital("c24", 24, 75_000.0),
                _capital("c200", 200, 400_000.0),
            ),
            owner_expense_items=(
                _expense("e1", 40_000.0, 1, 2),
                _expense("e2", 15_000.0, 2, None),
                _expense("e9", 99_000.0, 9, None),
            ),
        ),
    ),
    GoldenCase(
        "multiple_sign_changes",
        _inputs(),
        BusinessPlan(capital_items=(_capital("big", 30, 6_000_000.0),)),
        levered_status=IrrStatus.MULTIPLE_SIGN_CHANGES,
        unlevered_status=IrrStatus.MULTIPLE_SIGN_CHANGES,
    ),
    GoldenCase(
        "no_positive_cash_flow",
        _inputs(current_noi=300_000.0, exit_cap_rate=2.0, ltv=0.8, interest_rate=0.06),
        levered_status=IrrStatus.NO_POSITIVE_CASH_FLOW,
    ),
    GoldenCase(
        "very_high_return",
        _inputs(purchase_price=1_000_000.0, current_noi=900_000.0, ltv=0.5, hold_period=3, exit_cap_rate=0.05),
    ),
    GoldenCase(
        "near_total_loss_one_year",
        _inputs(current_noi=1.0, ltv=0.0, hold_period=1, exit_cap_rate=0.05),
    ),
    GoldenCase(
        "near_total_loss_three_years",
        _inputs(current_noi=1.0, ltv=0.0, hold_period=3, exit_cap_rate=0.05),
    ),
    GoldenCase(
        "long_hold",
        _inputs(hold_period=30, amortization=25, noi_growth=0.025, io_period=3),
    ),
    GoldenCase(
        "outside_anchor_search_domain",
        _inputs(purchase_price=1e12, current_noi=1e-4, ltv=0.0, hold_period=1, exit_cap_rate=0.05),
        levered_status=IrrStatus.ROOT_OUTSIDE_SEARCH_DOMAIN,
        unlevered_status=IrrStatus.ROOT_OUTSIDE_SEARCH_DOMAIN,
        expected_failures=frozenset(
            {
                "Levered IRR availability",
                "Levered IRR",
                "Unlevered IRR availability",
                "Unlevered IRR",
            }
        ),
    ),
)

CASES_BY_NAME = {case.name: case for case in GOLDEN_CASES}


def analyze(case: GoldenCase) -> AcquisitionResults:
    results = analyze_quick_acquisition_with_business_plan(case.inputs, business_plan=case.business_plan)
    assert results.levered_irr_status is case.levered_status, (case.name, results.levered_irr_status)
    assert results.unlevered_irr_status is case.unlevered_status, (case.name, results.unlevered_irr_status)
    return results


def source_for(case: GoldenCase, *, deal_name: str | None = None) -> QuickAuditSource:
    return QuickAuditSource(
        deal_id=f"golden-{case.name}",
        deal_name=deal_name if deal_name is not None else f"Golden {case.name}",
        asset_type_label="Multifamily",
        asset_subtype="Garden",
        inputs=case.inputs,
        business_plan=case.business_plan,
        results=analyze(case),
        analysis_fingerprint=fingerprint_quick_inputs(case.inputs, business_plan=case.business_plan),
        generated_at=GENERATED_AT,
        anchor_version="test",
        source_commit=None,
    )


def build(case: GoldenCase, **kwargs: str) -> bytes:
    return build_quick_audit_workbook(source_for(case, **kwargs))


# --- Reading workbooks -------------------------------------------------------


def load(data: bytes, *, values: bool = False) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(data), data_only=values)


def row_of(ws, label: str, *, column: int = 1, start: int = 1) -> int:  # noqa: ANN001
    """The first row at or after ``start`` whose ``column`` holds ``label``."""

    for row in range(start, ws.max_row + 1):
        if ws.cell(row, column).value == label:
            return row
    raise KeyError(f"{ws.title}: {label!r} not found")


def check_rows(values_wb: openpyxl.Workbook) -> dict[str, tuple[object, object, object, str]]:
    """Metric -> (Anchor, Excel, tolerance, status) for every reconciliation row."""

    ws = values_wb["Checks"]
    header = row_of(ws, "Metric")
    rows: dict[str, tuple[object, object, object, str]] = {}
    for row in range(header + 1, ws.max_row + 1):
        status = ws.cell(row, 6).value
        metric = ws.cell(row, 1).value
        if status is None or metric is None:
            continue
        rows[str(metric)] = (ws.cell(row, 2).value, ws.cell(row, 3).value, ws.cell(row, 5).value, str(status))
    return rows


def status_block(values_wb: openpyxl.Workbook) -> dict[str, object]:
    ws = values_wb["Checks"]
    block: dict[str, object] = {}
    for row in range(1, row_of(ws, "Reconciliation") + 1):
        label = ws.cell(row, 1).value
        if isinstance(label, str):
            block[label] = ws.cell(row, 2).value
    return block


# --- Mutating workbook XML (test-only) --------------------------------------


def _sheet_part(name: str) -> str:
    # XlsxWriter writes sheetN.xml in creation order, which is SHEET_ORDER.
    return f"xl/worksheets/sheet{SHEET_ORDER.index(name) + 1}.xml"


def _rewrite(data: bytes, sheet: str, edit) -> bytes:  # noqa: ANN001
    part = _sheet_part(sheet)
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == part:
                content = edit(content.decode("utf-8")).encode("utf-8")
            target.writestr(info, content)
    return out.getvalue()


def _cell_pattern(ref: str) -> re.Pattern[str]:
    return re.compile(rf'<c r="{ref}"(?P<attrs>[^>]*?)(?:/>|>(?P<body>.*?)</c>)', re.S)


def read_formula(data: bytes, sheet: str, ref: str) -> str:
    xml = zipfile.ZipFile(io.BytesIO(data)).read(_sheet_part(sheet)).decode("utf-8")
    match = _cell_pattern(ref).search(xml)
    assert match is not None and match.group("body") is not None, (sheet, ref)
    formula = re.search(r"<f>(.*?)</f>", match.group("body"), re.S)
    assert formula is not None, (sheet, ref)
    return formula.group(1).replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&amp;", "&")


def replace_formula(data: bytes, sheet: str, ref: str, new_formula: str) -> bytes:
    def edit(xml: str) -> str:
        match = _cell_pattern(ref).search(xml)
        assert match is not None, (sheet, ref)
        replacement = f'<c r="{ref}"{match.group("attrs")}><f>{escape(new_formula.lstrip("="))}</f><v></v></c>'
        return xml[: match.start()] + replacement + xml[match.end() :]

    return _rewrite(data, sheet, edit)


def blank_cell(data: bytes, sheet: str, ref: str) -> bytes:
    def edit(xml: str) -> str:
        match = _cell_pattern(ref).search(xml)
        assert match is not None, (sheet, ref)
        attrs = re.sub(r'\s+t="[^"]*"', "", match.group("attrs"))
        return xml[: match.start()] + f'<c r="{ref}"{attrs}/>' + xml[match.end() :]

    return _rewrite(data, sheet, edit)


def set_number(data: bytes, sheet: str, ref: str, value: float) -> bytes:
    def edit(xml: str) -> str:
        match = _cell_pattern(ref).search(xml)
        assert match is not None, (sheet, ref)
        attrs = re.sub(r'\s+t="[^"]*"', "", match.group("attrs"))
        return xml[: match.start()] + f'<c r="{ref}"{attrs}><v>{value!r}</v></c>' + xml[match.end() :]

    return _rewrite(data, sheet, edit)
