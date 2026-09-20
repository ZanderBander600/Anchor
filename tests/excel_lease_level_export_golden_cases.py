"""Excel Export 3 test support: golden Lease-Level cases and helpers.

Each case is analysed by Anchor's own Lease-Level entry point -- exactly what
the export re-runs -- so a case cannot silently stop exercising the rent-roll
behaviour it is named for.

Unlike Quick and Detailed there is no saved analysis to stand in for: a
Lease-Level Deal persists none, so ``source_for`` runs the authoritative
analysis and hands the result to the workbook, which is precisely what the
server route does.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timezone

import openpyxl

from anchor.analysis.business_plan_analysis import (
    analyze_lease_level_acquisition_with_business_plan,
)
from anchor.business_plan import (
    BusinessPlan,
    CapitalItemCategory,
    CapitalPlanItem,
    OwnerExpenseCategory,
    OwnerExpenseItem,
)
from anchor.contracts import AcquisitionTerms
from anchor.exports.excel import (
    LEASE_LEVEL_SHEETS as LEASE_LEVEL_SHEET_ORDER,
    LeaseLevelAuditSource,
    build_lease_level_audit_workbook,
)
from anchor.leasing import (
    EscalationBasis,
    InitialVacancyAssumptions,
    InitialVacancyStrategy,
    Lease,
    LeaseLevelOperatingInputs,
    LeaseLevelPropertyInputs,
    LeaseType,
    LeasingCommissionMethod,
    MarketLeasingAssumptions,
    RecoveryBasis,
    Suite,
)

GENERATED_AT = datetime(2026, 9, 19, 15, 30, 0, tzinfo=timezone.utc)

ANALYSIS_START = date(2027, 1, 1)
AREA = 100_000.0
PRICE = 40_000_000.0


# =============================================================================
# Raw input builders
# =============================================================================


def terms(**overrides: object) -> AcquisitionTerms:
    base: dict[str, object] = {
        "purchase_price": PRICE,
        "hold_period": 5,
        "exit_cap_rate": 0.065,
        "ltv": 0.60,
        "interest_rate": 0.05,
        "amortization": 30,
        "acquisition_cost_pct": 0.02,
        "financing_fee_pct": 0.01,
        "disposition_cost_pct": 0.015,
        "annual_capex_reserve": 0.0,
        "io_period": 0,
    }
    base.update(overrides)
    return AcquisitionTerms(**base)  # type: ignore[arg-type]


def property_inputs(area: float = AREA) -> LeaseLevelPropertyInputs:
    return LeaseLevelPropertyInputs(
        analysis_start_date=ANALYSIS_START, rentable_area_sf=area
    )


def market(**overrides: object) -> MarketLeasingAssumptions:
    base: dict[str, object] = {
        "market_rent_psf": 30.0,
        "market_rent_growth": 0.03,
        "renewal_rent_psf": None,
        "renewal_rent_spread": 0.0,
        "renewal_term_months": 60,
        "successor_escalation_pct": 0.0,
        "renewal_downtime_months": 0.0,
        "renewal_free_rent_months": 0.0,
        "new_term_months": 60,
        "new_downtime_months": 0.0,
        "new_free_rent_months": 0.0,
        "renewal_ti_psf": 0.0,
        "new_ti_psf": 0.0,
        "leasing_commission_method": (
            LeasingCommissionMethod.PCT_OF_TOTAL_CONTRACTUAL_BASE_RENT
        ),
        "renewal_lc_pct": 0.0,
        "new_lc_pct": 0.0,
        "renewal_probability": 0.7,
        "renewal_lease_type": LeaseType.NNN,
        "renewal_recovery_basis": None,
        "renewal_expense_stop_psf": None,
        "new_lease_type": LeaseType.NNN,
        "new_recovery_basis": None,
        "new_expense_stop_psf": None,
    }
    base.update(overrides)
    return MarketLeasingAssumptions(**base)  # type: ignore[arg-type]


def operating(**overrides: object) -> LeaseLevelOperatingInputs:
    base: dict[str, object] = {
        "other_income": 120_000.0,
        "other_income_growth": 0.03,
        "credit_loss_pct": 0.0,
        "property_taxes": 600_000.0,
        "insurance": 120_000.0,
        "utilities": 240_000.0,
        "repairs_maintenance": 180_000.0,
        "other_operating_expenses": 60_000.0,
        "management_fee_pct": 0.03,
        "expense_growth": 0.03,
        "recoverable_expense_ratio": 1.0,
    }
    base.update(overrides)
    return LeaseLevelOperatingInputs(**base)  # type: ignore[arg-type]


def suite(suite_id: str, area: float, **overrides: object) -> Suite:
    base: dict[str, object] = {"suite_id": suite_id, "suite_area_sf": area}
    base.update(overrides)
    return Suite(**base)  # type: ignore[arg-type]


def vacant(
    suite_id: str,
    area: float,
    *,
    strategy: InitialVacancyStrategy = InitialVacancyStrategy.MARKET_LEASE_UP,
    lease_up: float | None = 0.0,
    **overrides: object,
) -> Suite:
    return suite(
        suite_id,
        area,
        initial_vacancy=InitialVacancyAssumptions(
            strategy=strategy, initial_lease_up_months=lease_up
        ),
        **overrides,
    )


def lease(
    for_suite: Suite,
    *,
    end: date = date(2029, 12, 31),
    start: date = date(2025, 1, 1),
    base_rent_psf: float = 30.0,
    escalation_pct: float = 0.0,
    escalation_basis: EscalationBasis = EscalationBasis.NONE,
    lease_type: LeaseType = LeaseType.NNN,
    recovery_basis: RecoveryBasis | None = None,
    expense_stop_psf: float | None = None,
    tenant_name: str | None = None,
    lease_id: str | None = None,
) -> Lease:
    return Lease(
        lease_id=lease_id or f"L-{for_suite.suite_id}",
        suite_id=for_suite.suite_id,
        tenant_name=tenant_name,
        leased_area_sf=for_suite.suite_area_sf,
        rent_commencement_date=start,
        lease_expiration_date=end,
        base_rent_psf=base_rent_psf,
        escalation_pct=escalation_pct,
        escalation_basis=escalation_basis,
        lease_type=lease_type,
        recovery_basis=recovery_basis,
        expense_stop_psf=expense_stop_psf,
    )


def capital(item_id: str, month: int, amount: float) -> CapitalPlanItem:
    return CapitalPlanItem(
        item_id=item_id,
        description=f"Capital {item_id}",
        category=CapitalItemCategory.VALUE_ADD_RENOVATION,
        month=month,
        amount=amount,
    )


def owner_expense(item_id: str, amount: float, first: int, last: int | None) -> OwnerExpenseItem:
    return OwnerExpenseItem(
        item_id=item_id,
        description=f"Expense {item_id}",
        category=OwnerExpenseCategory.ASSET_MANAGEMENT,
        annual_amount=amount,
        first_year=first,
        last_year=last,
    )


# =============================================================================
# Golden cases
# =============================================================================


@dataclass(frozen=True)
class LeaseLevelGoldenCase:
    """One deterministic Lease-Level deal, named for what it exercises."""

    name: str
    suites: tuple[Suite, ...]
    leases: tuple[Lease, ...]
    market: MarketLeasingAssumptions
    operating: LeaseLevelOperatingInputs
    terms: AcquisitionTerms = field(default_factory=terms)
    property_inputs: LeaseLevelPropertyInputs = field(default_factory=property_inputs)
    business_plan: BusinessPlan = field(default_factory=BusinessPlan)
    deal_name: str = "Rivermark Center"


_S1 = suite("S1", 60_000.0, suite_label="Suite 100")
_S2 = suite("S2", 40_000.0, suite_label="Suite 200")
_WHOLE = suite("S1", AREA, suite_label="Whole building")

#: Analyst-authored identifiers that Excel would read as formulas if they
#: were ever written as anything but literal text.
_HOSTILE_A = suite("=cmd|S1", 60_000.0, suite_label='=HYPERLINK("http://x")')
_HOSTILE_B = suite("@S2", 40_000.0, suite_label="+1-2")

#: A suite whose full override is genuinely different from the property
#: default in every dimension the model reads.
_OVERRIDE = market(
    market_rent_psf=42.0,
    market_rent_growth=0.01,
    renewal_probability=0.25,
    renewal_rent_spread=-0.10,
    renewal_term_months=36,
    new_term_months=36,
    new_downtime_months=6.0,
    renewal_free_rent_months=1.0,
    new_free_rent_months=3.0,
    renewal_ti_psf=10.0,
    new_ti_psf=40.0,
    renewal_lc_pct=0.03,
    new_lc_pct=0.06,
    successor_escalation_pct=0.025,
)


GOLDEN_CASES: tuple[LeaseLevelGoldenCase, ...] = (
    LeaseLevelGoldenCase(
        name="occupied_no_in_hold_rollover",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2045, 12, 31)),),
        market=market(),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="multiple_suites",
        suites=(_S1, _S2),
        leases=(lease(_S1, end=date(2029, 12, 31)), lease(_S2, end=date(2031, 6, 30))),
        market=market(),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="fractional_downtime_boundary",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2028, 6, 30)),),
        market=market(new_downtime_months=2.25, renewal_downtime_months=0.5),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="free_rent_both_branches",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2028, 12, 31)),),
        market=market(renewal_free_rent_months=2.0, new_free_rent_months=6.0),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="contractual_rent_steps",
        suites=(_WHOLE,),
        leases=(
            lease(
                _WHOLE, end=date(2032, 12, 31), escalation_pct=0.03,
                escalation_basis=EscalationBasis.LEASE_ANNIVERSARY,
            ),
        ),
        market=market(successor_escalation_pct=0.025),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="renewal_probability_zero",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2028, 12, 31)),),
        market=market(renewal_probability=0.0, new_downtime_months=4.0, new_ti_psf=30.0),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="renewal_probability_one",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2028, 12, 31)),),
        market=market(renewal_probability=1.0, renewal_ti_psf=8.0, renewal_lc_pct=0.02),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="renewal_probability_blended_with_downtime",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2028, 12, 31)),),
        market=market(
            renewal_probability=0.55, new_downtime_months=9.0,
            renewal_rent_spread=-0.05, new_ti_psf=35.0, new_lc_pct=0.05,
            renewal_ti_psf=7.0, renewal_lc_pct=0.02,
        ),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="suite_level_override",
        suites=(_S1, suite("S2", 40_000.0, suite_label="Suite 200", market_leasing_override=_OVERRIDE)),
        leases=(lease(_S1, end=date(2029, 12, 31)), lease(_S2, end=date(2028, 9, 30))),
        market=market(),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="suite_rent_level_override_only",
        suites=(_S1, suite("S2", 40_000.0, suite_label="Suite 200", market_rent_psf=55.0)),
        leases=(lease(_S1, end=date(2029, 12, 31)), lease(_S2, end=date(2028, 9, 30))),
        market=market(),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="initial_vacancy_hold_vacant",
        suites=(
            _S1,
            vacant("S2", 40_000.0, strategy=InitialVacancyStrategy.HOLD_VACANT, lease_up=None, suite_label="Suite 200"),
        ),
        leases=(lease(_S1, end=date(2031, 12, 31)),),
        market=market(),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="initial_vacancy_market_lease_up",
        suites=(_S1, vacant("S2", 40_000.0, lease_up=7.5, suite_label="Suite 200")),
        leases=(lease(_S1, end=date(2031, 12, 31)),),
        market=market(new_ti_psf=30.0, new_lc_pct=0.05, new_free_rent_months=3.0),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="recovery_structures_all_three",
        suites=(
            suite("S1", 40_000.0, suite_label="NNN"),
            suite("S2", 35_000.0, suite_label="Gross"),
            suite("S3", 25_000.0, suite_label="Modified gross"),
        ),
        leases=(
            lease(suite("S1", 40_000.0), end=date(2031, 12, 31), lease_type=LeaseType.NNN),
            lease(suite("S2", 35_000.0), end=date(2031, 12, 31), lease_type=LeaseType.GROSS),
            lease(
                suite("S3", 25_000.0), end=date(2031, 12, 31),
                lease_type=LeaseType.MODIFIED_GROSS,
                recovery_basis=RecoveryBasis.EXPENSE_STOP_PSF,
                expense_stop_psf=6.0,
            ),
        ),
        market=market(),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="successor_modified_gross_with_stop",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2028, 12, 31)),),
        market=market(
            renewal_lease_type=LeaseType.MODIFIED_GROSS,
            renewal_recovery_basis=RecoveryBasis.EXPENSE_STOP_PSF,
            renewal_expense_stop_psf=5.0,
            new_lease_type=LeaseType.GROSS,
        ),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="partly_recoverable_expenses",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2031, 12, 31)),),
        market=market(),
        operating=operating(recoverable_expense_ratio=0.55),
    ),
    LeaseLevelGoldenCase(
        name="credit_loss_and_management_fee",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2031, 12, 31)),),
        market=market(),
        operating=operating(credit_loss_pct=0.04, management_fee_pct=0.05),
    ),
    LeaseLevelGoldenCase(
        name="ti_lc_inside_hold_and_forward_window",
        suites=(_S1, _S2),
        leases=(
            # S1 rolls inside the hold; S2 rolls inside the forward window.
            lease(_S1, end=date(2029, 6, 30)),
            lease(_S2, end=date(2032, 3, 31)),
        ),
        market=market(
            renewal_ti_psf=12.0, new_ti_psf=45.0,
            renewal_lc_pct=0.03, new_lc_pct=0.06, new_downtime_months=5.0,
        ),
        operating=operating(),
    ),
    LeaseLevelGoldenCase(
        name="growth_divergence",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2031, 12, 31)),),
        market=market(market_rent_growth=0.06),
        operating=operating(expense_growth=-0.01, other_income_growth=0.08),
    ),
    LeaseLevelGoldenCase(
        name="business_plan_across_the_hold",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2031, 12, 31)),),
        market=market(),
        operating=operating(),
        business_plan=BusinessPlan(
            capital_items=(
                capital("closing", 0, 750_000.0),
                capital("mid", 18, 400_000.0),
                capital("post", 72, 250_000.0),
            ),
            owner_expense_items=(owner_expense("am", 60_000.0, 1, None),),
        ),
    ),
    LeaseLevelGoldenCase(
        name="interest_only_debt",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2031, 12, 31)),),
        market=market(),
        operating=operating(),
        terms=terms(io_period=10),
    ),
    LeaseLevelGoldenCase(
        name="all_cash",
        suites=(_WHOLE,),
        leases=(lease(_WHOLE, end=date(2031, 12, 31)),),
        market=market(),
        operating=operating(),
        terms=terms(ltv=0.0, interest_rate=0.0, financing_fee_pct=0.0),
    ),
    LeaseLevelGoldenCase(
        name="ten_year_hold_many_generations",
        suites=(_S1, _S2),
        leases=(lease(_S1, end=date(2028, 12, 31)), lease(_S2, end=date(2027, 12, 31))),
        market=market(
            renewal_term_months=24, new_term_months=36, new_downtime_months=2.5,
            renewal_free_rent_months=1.0, new_free_rent_months=4.0,
            renewal_ti_psf=6.0, new_ti_psf=28.0, renewal_lc_pct=0.02, new_lc_pct=0.05,
            successor_escalation_pct=0.03,
        ),
        operating=operating(),
        terms=terms(hold_period=10),
    ),
    LeaseLevelGoldenCase(
        name="mutation_probe",
        suites=(_S1, _S2),
        leases=(
            # S1 rolls early inside the hold; S2 rolls inside the forward window.
            lease(_S1, end=date(2028, 6, 30)),
            lease(_S2, end=date(2032, 3, 31)),
        ),
        market=market(
            # Every line a mutation targets carries a real value here: free
            # rent on both branches, TI and LC on both, a renewal spread that
            # makes the two branches genuinely price differently, and downtime
            # that moves the new tenant into a later market band.
            renewal_probability=0.6,
            renewal_rent_spread=-0.08,
            renewal_free_rent_months=3.0,
            new_free_rent_months=6.0,
            renewal_downtime_months=1.5,
            new_downtime_months=7.0,
            renewal_ti_psf=12.0,
            new_ti_psf=45.0,
            renewal_lc_pct=0.03,
            new_lc_pct=0.06,
            successor_escalation_pct=0.025,
        ),
        operating=operating(credit_loss_pct=0.03, recoverable_expense_ratio=0.8),
    ),
    LeaseLevelGoldenCase(
        name="hostile_text_values",
        suites=(_HOSTILE_A, _HOSTILE_B),
        leases=(
            lease(
                _HOSTILE_A, end=date(2031, 12, 31),
                lease_id="@SUM(A1:A9)", tenant_name="-1+1",
            ),
            lease(
                _HOSTILE_B, end=date(2031, 12, 31),
                lease_id="+CMD", tenant_name='=HYPERLINK("http://x","click")',
            ),
        ),
        market=market(),
        operating=operating(),
        deal_name='=cmd|" /C calc"!A0',
    ),
)

CASES_BY_NAME = {case.name: case for case in GOLDEN_CASES}


# =============================================================================
# Building
# =============================================================================


def analyze(case: LeaseLevelGoldenCase):
    """Run Anchor's own Lease-Level entry point over the case's inputs."""

    return analyze_lease_level_acquisition_with_business_plan(
        case.terms,
        case.property_inputs,
        case.suites,
        case.leases,
        market_leasing=case.market,
        operating_inputs=case.operating,
        business_plan=case.business_plan,
    )


def source_for(case: LeaseLevelGoldenCase, *, deal_name: str | None = None) -> LeaseLevelAuditSource:
    analysis = analyze(case)
    return LeaseLevelAuditSource(
        deal_id=f"deal-{case.name}",
        deal_name=deal_name if deal_name is not None else case.deal_name,
        asset_type_label="Office",
        asset_subtype=None,
        terms=case.terms,
        property_inputs=case.property_inputs,
        operating_inputs=case.operating,
        market_leasing=case.market,
        suites=case.suites,
        leases=case.leases,
        business_plan=case.business_plan,
        monthly=analysis.monthly_projection,
        annual=analysis.annual_projection,
        results=analysis.results,
        analysis_fingerprint=f"fingerprint-{case.name}",
        generated_at=GENERATED_AT,
        anchor_version="0.0.0-test",
        source_commit=None,
    )


def build(case: LeaseLevelGoldenCase, **kwargs: str) -> bytes:
    return build_lease_level_audit_workbook(source_for(case, **kwargs))


def load(data: bytes, *, values: bool = False) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(data), data_only=values)


# =============================================================================
# Reading a built workbook, without opening Excel
# =============================================================================


def row_of(ws, label: str, *, column: int = 1, start: int = 1) -> int:  # noqa: ANN001
    for row in range(start, ws.max_row + 1):
        if ws.cell(row, column).value == label:
            return row
    raise AssertionError(f"no row labelled {label!r}")


def check_rows(values_wb: openpyxl.Workbook) -> dict[str, tuple[object, object, object, str]]:
    """Every Checks row, by metric: ``(anchor, excel, difference, status)``."""

    ws = values_wb["Checks"]
    header = row_of(ws, "Metric")
    rows: dict[str, tuple[object, object, object, str]] = {}
    for row in range(header + 1, ws.max_row + 1):
        metric = ws.cell(row, 1).value
        status = ws.cell(row, 6).value
        if metric is None or status is None:
            continue
        rows[str(metric)] = (
            ws.cell(row, 2).value, ws.cell(row, 3).value, ws.cell(row, 4).value, str(status)
        )
    return rows


def status_block(values_wb: openpyxl.Workbook) -> dict[str, object]:
    ws = values_wb["Checks"]
    block: dict[str, object] = {}
    for row in range(1, 20):
        label = ws.cell(row, 1).value
        if label is None:
            continue
        block[str(label)] = ws.cell(row, 2).value
    return block


def _sheet_part(name: str) -> str:
    """The zip entry holding one worksheet's XML, by its display name."""

    index = list(LEASE_LEVEL_SHEET_ORDER).index(name) + 1
    return f"xl/worksheets/sheet{index}.xml"



def _rewrite(data: bytes, sheet: str, edit) -> bytes:  # noqa: ANN001
    """Return ``data`` with one worksheet's XML transformed by ``edit``."""

    part = _sheet_part(sheet)
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename == part:
                payload = edit(payload.decode("utf-8")).encode("utf-8")
            target.writestr(item, payload)
    return output.getvalue()


def _cell_pattern(ref: str) -> re.Pattern[str]:
    return re.compile(rf'(<c r="{ref}"[^>]*>)(.*?)(</c>)', re.DOTALL)


def read_formula(data: bytes, sheet: str, ref: str) -> str:
    """The formula text stored in one cell."""

    xml = zipfile.ZipFile(io.BytesIO(data)).read(_sheet_part(sheet)).decode("utf-8")
    match = _cell_pattern(ref).search(xml)
    if match is None:
        raise AssertionError(f"no cell {ref} on {sheet}")
    formula = re.search(r"<f[^>]*>(.*?)</f>", match.group(2), re.DOTALL)
    if formula is None:
        raise AssertionError(f"{sheet}!{ref} holds no formula")
    return formula.group(1)


def replace_formula(data: bytes, sheet: str, ref: str, new_formula: str) -> bytes:
    """Replace one cell's formula, leaving every other cell untouched."""

    def edit(xml: str) -> str:
        def swap(match: re.Match[str]) -> str:
            body = re.sub(r"<f[^>]*>.*?</f>", f"<f>{new_formula}</f>", match.group(2), flags=re.DOTALL)
            return match.group(1) + body + match.group(3)

        replaced, count = _cell_pattern(ref).subn(swap, xml)
        if count != 1:
            raise AssertionError(f"expected one {ref} cell on {sheet}, found {count}")
        return replaced

    return _rewrite(data, sheet, edit)


def blank_cell(data: bytes, sheet: str, ref: str) -> bytes:
    """Delete one cell's formula entirely, as a careless edit would."""

    def edit(xml: str) -> str:
        replaced, count = _cell_pattern(ref).subn(lambda m: m.group(1) + m.group(3), xml)
        if count != 1:
            raise AssertionError(f"expected one {ref} cell on {sheet}, found {count}")
        return replaced

    return _rewrite(data, sheet, edit)


def set_number(data: bytes, sheet: str, ref: str, value: float) -> bytes:
    """Replace one cell's contents with a literal number, as an analyst
    editing a Working Input would."""

    def edit(xml: str) -> str:
        def swap(match: re.Match[str]) -> str:
            opening = re.sub(r'\st="[^"]*"', "", match.group(1))
            return opening + f"<v>{value!r}</v>" + match.group(3)

        replaced, count = _cell_pattern(ref).subn(swap, xml)
        if count != 1:
            raise AssertionError(f"expected one {ref} cell on {sheet}, found {count}")
        return replaced

    return _rewrite(data, sheet, edit)
