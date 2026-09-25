"""Refinance & Capital Events V1 Stage 3 -- product presentation facts.

``docs/architecture/REFINANCE_CAPITAL_EVENTS_V1.md`` Section 25.1:

1. **Presence.** ``GET /investments/{id}/capital-event-presence`` states which
   Strategies configure a refinance and which Project metrics are then the
   acquisition-financing reference. The frontend never keeps its own list.
2. **The implicit Common Equity perspective.** A structure with a refinance and
   no Common Equity marker still gets a Common Equity matrix, because the
   refinance-adjusted return lives only there.
3. **Unexecuted cells** carry a typed reason and an analyst sentence, never a
   figure.
4. **One vocabulary.** The Python and TypeScript reason catalogs cover the
   same typed reasons; the reference label is spelled identically everywhere.
5. **Exports 1-3** label their levered figures as the acquisition-financing
   reference exactly when a refinance is configured, and are otherwise
   unchanged.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import _refinance_v1_stage_3_fixtures as s3  # type: ignore[import-not-found]
import _refinance_v1_stage_2_fixtures as fx  # type: ignore[import-not-found]
from anchor.capital_structure.refinance_contracts import RefinanceUnavailableReason
from anchor.deals import refinance_presentation as presentation
from anchor.deals import store
from anchor.deals.structured_variants import IMPLICIT_COMMON_EQUITY_ID
from anchor.exports.excel import _workbook

ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / "web" / "src" / "refinanceCatalog.ts"
PRESENCE_METRICS = ["levered_irr", "equity_multiple", "total_profit", "total_equity_invested", "min_dscr"]


@pytest.fixture
def db(tmp_path: Path) -> Path:
    return tmp_path / "anchor.db"


@pytest.fixture
def client(db: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    return s3.client(db, monkeypatch)


# =============================================================================
# 1. Presence
# =============================================================================


def test_presence_names_a_refinance_bearing_strategy_and_the_reference_metrics(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db)
    body = client.get(f"/investments/{investment_id}/capital-event-presence").json()

    assert body["strategies"] == [{"strategy_id": "base", "capital_events_configured": True}]
    assert body["acquisition_financing_metrics"] == PRESENCE_METRICS


def test_presence_is_false_without_a_refinance_and_the_metrics_stay_listed(db: Path, client: TestClient) -> None:
    deal = fx.base_deal(db, name="Plain")
    investment_id, _ = store.set_deal_capital_structure(deal.id, fx.closing_mezz_only(deal.id), db_path=db)
    body = client.get(f"/investments/{investment_id}/capital-event-presence").json()

    assert body["strategies"] == [{"strategy_id": "base", "capital_events_configured": False}]


def test_presence_of_an_unknown_investment_is_not_found(client: TestClient) -> None:
    assert client.get("/investments/nope/capital-event-presence").status_code == 404


def test_the_project_matrix_is_unchanged_by_a_refinance(db: Path, client: TestClient) -> None:
    """The Project namespace keeps the acquisition-loan figures (R-P rule 3):
    the matrix is not rewritten; presence tells the product how to label it."""

    _, investment_id = s3.refinance_deal(db)
    cell = client.post(f"/investments/{investment_id}/decision-matrix", json={}).json()["matrix"]["cells"][0]

    assert cell["results"]["levered_cash_flows"][0] == -4_000_000.0
    assert cell["results"]["remaining_loan_balance"] == 4_800_000.0


# =============================================================================
# 2-3. The Common Equity and replacement position matrices
# =============================================================================


def _position_cell(client: TestClient, investment_id: str, position_id: str) -> dict[str, Any]:
    response = client.post(f"/investments/{investment_id}/position-decision-matrix/{position_id}")
    assert response.status_code == 200, response.text
    (cell,) = response.json()["matrix"]["cells"]
    return cell


def _metrics(cell: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {item["metric"]: item for item in cell["metrics"]}


def test_the_implicit_common_equity_matrix_reports_the_refinance_adjusted_return(
    db: Path, client: TestClient
) -> None:
    _, investment_id = s3.refinance_deal(db)
    metrics = _metrics(_position_cell(client, investment_id, IMPLICIT_COMMON_EQUITY_ID))

    # 560 + 2,920 + 400 + 400 + 4,900 thousand returned on 4,000 invested.
    assert metrics["total_cash_returned"]["value"] == pytest.approx(9_180_000.0)
    assert metrics["equity_multiple"]["value"] == pytest.approx(2.295)
    assert metrics["common_equity_irr"]["irr_status"] == "defined"


def test_an_unexecuted_refinance_blanks_common_equity_with_a_typed_reason(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db, s3.f4_structure)
    cell = _position_cell(client, investment_id, IMPLICIT_COMMON_EQUITY_ID)

    assert cell["applicability"] == "present"
    for metric in cell["metrics"]:
        assert metric["value"] is None
        assert metric["reason"] == "refinance_unavailable"
        assert "“Year-2 refinance” did not execute" in metric["message"]


def test_an_unexecuted_replacement_is_a_present_cell_with_no_figure(db: Path, client: TestClient) -> None:
    _, investment_id = s3.refinance_deal(db, s3.f4_structure)
    cell = _position_cell(client, investment_id, fx.REPLACEMENT_ID)

    assert cell["applicability"] == "present"
    assert cell["unavailable_message"]
    assert fx.REPLACEMENT_ID not in cell["unavailable_message"]
    assert all(metric["value"] is None for metric in cell["metrics"])


def test_no_implicit_common_equity_without_a_refinance(db: Path, client: TestClient) -> None:
    deal = fx.base_deal(db, name="Plain")
    investment_id, _ = store.set_deal_capital_structure(deal.id, fx.closing_mezz_only(deal.id), db_path=db)
    response = client.post(f"/investments/{investment_id}/position-decision-matrix/{IMPLICIT_COMMON_EQUITY_ID}")

    assert response.status_code == 404


# =============================================================================
# 4. One vocabulary
# =============================================================================


def _typescript_reason_keys() -> set[str]:
    source = CATALOG.read_text(encoding="utf-8")
    block = source.split("const REASON_SENTENCES", 1)[1].split("\n};", 1)[0]
    return set(re.findall(r"^  ([a-z_]+): ", block, flags=re.MULTILINE))


def test_python_and_typescript_word_every_typed_reason() -> None:
    reasons = {reason.value for reason in RefinanceUnavailableReason}

    assert {reason.value for reason in presentation.REFINANCE_REASON_SENTENCES} == reasons
    assert _typescript_reason_keys() == reasons


def test_the_reference_label_is_spelled_once_everywhere() -> None:
    source = CATALOG.read_text(encoding="utf-8")

    assert presentation.ACQUISITION_REFERENCE_LABEL == s3.REFERENCE
    assert _workbook.ACQUISITION_REFERENCE_LABEL == s3.REFERENCE
    assert f"ACQUISITION_REFERENCE_LABEL = '{s3.REFERENCE}'" in source


def test_no_reason_sentence_quotes_an_identity_or_a_zero_fallback() -> None:
    for sentence in presentation.REFINANCE_REASON_SENTENCES.values():
        assert "_id" not in sentence
        assert "$0" not in sentence


# =============================================================================
# 5. Exports 1-3 label the acquisition-financing reference
# =============================================================================


def _summary_text(client: TestClient, deal_id: str) -> list[str]:
    response = client.get(f"/deals/{deal_id}/exports/quick-underwrite.xlsx")
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content))["Summary"]
    return [str(cell.value) for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str)]


def test_a_quick_export_labels_the_levered_figures_when_a_refinance_exists(db: Path, client: TestClient) -> None:
    deal, _ = s3.refinance_deal(db)
    s3.analysed(db, deal)
    text = _summary_text(client, deal.id)

    assert _workbook.REFINANCE_NOTICE in text
    assert any(s3.REFERENCE in value for value in text)


def test_a_quick_export_without_a_refinance_carries_no_reference_wording(db: Path, client: TestClient) -> None:
    deal = fx.base_deal(db, name="Plain")
    store.set_deal_capital_structure(deal.id, fx.closing_mezz_only(deal.id), db_path=db)
    s3.analysed(db, deal)
    text = _summary_text(client, deal.id)

    assert _workbook.REFINANCE_NOTICE not in text
    assert not any(s3.REFERENCE in value for value in text)


def test_a_refinance_changes_nothing_in_the_quick_export_but_its_labels(db: Path, client: TestClient) -> None:
    """Every number cell of the Summary is identical with and without the
    refinance: the export is labeled, never recomputed."""

    plain = fx.base_deal(db, name="Same")
    evented, _ = s3.refinance_deal(db, name="Same")
    s3.analysed(db, plain)
    s3.analysed(db, evented)

    def numbers(deal_id: str) -> list[Any]:
        response = client.get(f"/deals/{deal_id}/exports/quick-underwrite.xlsx")
        assert response.status_code == 200, response.text
        book = load_workbook(io.BytesIO(response.content))
        return [
            cell.value
            for name in book.sheetnames
            if name != "Audit Metadata"
            for row in book[name].iter_rows()
            for cell in row
            if isinstance(cell.value, (int, float))
        ]

    assert numbers(plain.id) == numbers(evented.id)
